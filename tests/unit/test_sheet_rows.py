"""_sheet_rows: 行解析只去末尾空行, 中间空行必须保留(行号偏移会导致模板落位全错)。"""
import io

from backend_v2.routes.bom_lists import _sheet_rows


def _make_xls_with_blank_middle():
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")
    for r, v in enumerate(["A", "", "C", "", "E"]):  # 中间空行在 r=1,3
        ws.write(r, 0, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xls_middle_blank_rows_kept():
    content = _make_xls_with_blank_middle()
    sheets = _sheet_rows("xls", content)
    rows = sheets[0]["rows"]
    assert len(rows) == 5, f"中间空行被丢弃导致行号偏移: {len(rows)}"
    assert rows[1][0] == "" and rows[3][0] == ""
    assert rows[0][0] == "A" and rows[2][0] == "C" and rows[4][0] == "E"


def test_xls_trailing_blank_rows_stripped():
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("S")
    ws.write(0, 0, "keep")
    for r in (1, 2, 3):
        ws.write(r, 0, "")
    buf = io.BytesIO()
    wb.save(buf)
    sheets = _sheet_rows("xls", buf.getvalue())
    assert len(sheets[0]["rows"]) == 1, "末尾空行应被剥离"


def test_xlsx_middle_blank_rows_kept():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, v in enumerate(["X", "", "Y"]):
        ws.cell(row=r + 1, column=1, value=v or None)
    buf = io.BytesIO()
    wb.save(buf)
    sheets = _sheet_rows("xlsx", buf.getvalue())
    rows = sheets[0]["rows"]
    assert len(rows) == 3, f"xlsx 中间空行被丢弃: {len(rows)}"


def test_empty_file_returns_empty_sheets():
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    sheets = _sheet_rows("xls", buf.getvalue())
    assert sheets[0]["rows"] == []
