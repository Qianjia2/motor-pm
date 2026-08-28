"""Excel export endpoints for projects, weekly reports, and risks."""
import io
from datetime import date
from flask import Blueprint, request, send_file, jsonify
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend.models import Project, WeeklyReport, RiskIssue, UserAuth

bp = Blueprint('export', __name__)


def check_token():
    """Allow token via query param (for window.open) or Authorization header."""
    token = request.args.get('token') or request.headers.get('Authorization', '').replace('Bearer ', '')
    if not token:
        return False
    user = UserAuth.query.filter_by(token=token, is_active=True).first()
    return user is not None

HEADER_FILL = PatternFill(start_color='1d4ed8', end_color='1d4ed8', fill_type='solid')
HEADER_FONT = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
BODY_FONT = Font(name='微软雅黑', size=10)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)

def style_header(ws, headers, row=1):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = THIN_BORDER

def style_body(ws, start_row, end_row, max_col):
    for r in range(start_row, end_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical='center')


@bp.route('/export/projects', methods=['GET'])
def export_projects():
    if not check_token():
        return jsonify({'error': '未登录'}), 401
    projects = Project.query.filter_by(is_active=True).order_by(Project.priority).all()

    wb = Workbook()
    ws = wb.active
    ws.title = '项目总表'

    headers = ['编号', '名称', '客户', '优先级', '当前阶段', '完成率%', '状态',
               '难度', '开始日期', '计划结束', '总预算', '已花费', '截止类型']
    style_header(ws, headers)

    for i, p in enumerate(projects, 2):
        ws.cell(row=i, column=1, value=p.code)
        ws.cell(row=i, column=2, value=p.name)
        ws.cell(row=i, column=3, value=p.client.name if p.client else '')
        ws.cell(row=i, column=4, value=p.priority)
        ws.cell(row=i, column=5, value=p.current_phase.name if p.current_phase else '')
        ws.cell(row=i, column=6, value=p.completion_pct)
        ws.cell(row=i, column=7, value=p.overall_status)
        ws.cell(row=i, column=8, value=p.difficulty)
        ws.cell(row=i, column=9, value=p.start_date.isoformat() if p.start_date else '')
        ws.cell(row=i, column=10, value=p.planned_end_date.isoformat() if p.planned_end_date else '')
        ws.cell(row=i, column=11, value=p.budget_total)
        ws.cell(row=i, column=12, value=p.budget_spent)
        ws.cell(row=i, column=13, value=p.deadline_type)

    style_body(ws, 2, len(projects) + 1, len(headers))
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'项目总表_{date.today().isoformat()}.xlsx')


@bp.route('/export/projects/<int:pid>/reports', methods=['GET'])
def export_reports(pid):
    if not check_token():
        return jsonify({'error': '未登录'}), 401
    project = Project.query.get_or_404(pid)
    reports = WeeklyReport.query.filter_by(project_id=pid).order_by(
        WeeklyReport.year.desc(), WeeklyReport.week_number.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = f'{project.code}_周报'

    headers = ['年度', '周次', '报告日期', '总体进展', '关键成果',
               '技术线', '本周进展', '下周计划', '状态', '风险/难点']
    style_header(ws, headers)

    row = 2
    for r in reports:
        for li in r.line_items:
            line_name = li.technical_line.short_name if li.technical_line else f'Line{li.line_id}'
            ws.cell(row=row, column=1, value=r.year)
            ws.cell(row=row, column=2, value=r.week_number)
            ws.cell(row=row, column=3, value=r.report_date.isoformat() if r.report_date else '')
            ws.cell(row=row, column=4, value=r.overall_progress)
            ws.cell(row=row, column=5, value=r.key_accomplishments)
            ws.cell(row=row, column=6, value=line_name)
            ws.cell(row=row, column=7, value=li.this_week)
            ws.cell(row=row, column=8, value=li.next_week)
            ws.cell(row=row, column=9, value=li.status)
            ws.cell(row=row, column=10, value=li.blocker)
            row += 1

    style_body(ws, 2, row - 1, len(headers))
    for col_letter, w in [('D', 30), ('E', 30), ('G', 35), ('H', 35), ('J', 25)]:
        ws.column_dimensions[col_letter].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'{project.code}_周报_{date.today().isoformat()}.xlsx')


@bp.route('/export/projects/<int:pid>/risks', methods=['GET'])
def export_risks(pid):
    if not check_token():
        return jsonify({'error': '未登录'}), 401
    project = Project.query.get_or_404(pid)
    risks = RiskIssue.query.filter_by(project_id=pid).order_by(RiskIssue.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = f'{project.code}_风险'

    headers = ['类型', '标题', '描述', '概率', '影响', '风险等级', '状态',
               '缓解措施', '负责人', '识别日期', '目标解决日期', '验证结果']
    style_header(ws, headers)

    for i, r in enumerate(risks, 2):
        ws.cell(row=i, column=1, value='风险' if r.type == 'risk' else '问题')
        ws.cell(row=i, column=2, value=r.title)
        ws.cell(row=i, column=3, value=r.description or '')
        ws.cell(row=i, column=4, value=r.probability or '')
        ws.cell(row=i, column=5, value=r.impact or '')
        ws.cell(row=i, column=6, value=r.risk_level or '')
        ws.cell(row=i, column=7, value=r.status or '')
        ws.cell(row=i, column=8, value=r.mitigation_plan or '')
        ws.cell(row=i, column=9, value=r.owner.name if r.owner else '')
        ws.cell(row=i, column=10, value=r.identified_date.isoformat() if r.identified_date else '')
        ws.cell(row=i, column=11, value=r.target_resolve_date.isoformat() if r.target_resolve_date else '')
        ws.cell(row=i, column=12, value=r.resolution or '')

    style_body(ws, 2, len(risks) + 1, len(headers))
    for col_letter, w in [('B', 30), ('C', 35), ('H', 30), ('L', 30)]:
        ws.column_dimensions[col_letter].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'{project.code}_风险_{date.today().isoformat()}.xlsx')
