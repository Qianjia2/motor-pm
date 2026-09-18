# -*- coding: utf-8 -*-
import io, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font(run, name_ascii='Times New Roman', name_cn='宋体', size=12, bold=False):
    run.font.name = name_ascii
    run.font.size = Pt(size)
    run.bold = bold
    r = run._element.rPr.rFonts
    r.set(qn('w:eastAsia'), name_cn)

def shade(cell, color='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def add_bold_runs(par, text, size=12, cn='宋体'):
    """解析 **粗体** 行内标记"""
    for i, seg in enumerate(re.split(r'(\*\*.+?\*\*)', text)):
        if not seg: continue
        b = seg.startswith('**') and seg.endswith('**')
        run = par.add_run(seg[2:-2] if b else seg)
        set_font(run, size=size, bold=b, name_cn=cn)

def convert(md_path, docx_path, title, subtitle, meta_lines):
    t = io.open(md_path, encoding='utf-8').read()
    doc = Document()
    # 页面设置
    for s in doc.sections:
        s.top_margin = Cm(2.54); s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(3.17); s.right_margin = Cm(3.17)

    # ===== 封面 =====
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(title), size=26, bold=True, name_cn='黑体')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(subtitle), size=14, name_cn='楷体')
    for _ in range(6): doc.add_paragraph()
    for line in meta_lines:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(line), size=12, name_cn='宋体')
    doc.add_page_break()

    # ===== 正文 =====
    lines = t.split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        # 表格
        if ln.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s\-:|]+\|$', lines[i+1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not re.match(r'^[\s\-:|]+$', '|'.join(cells)):
                    rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            tbl = doc.add_table(rows=0, cols=ncol)
            tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                cells = tbl.add_row().cells
                for ci in range(ncol):
                    txt = row[ci] if ci < len(row) else ''
                    cp = cells[ci].paragraphs[0]
                    add_bold_runs(cp, txt, size=9)
                    if ri == 0:
                        shade(cells[ci]); 
                        for r_ in cp.runs: r_.bold = True
            doc.add_paragraph()
            continue
        # 标题
        m = re.match(r'^(#{1,4})\s+(.+)$', ln)
        if m:
            lvl = len(m.group(1)); txt = m.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(6)
            sz = {1:18, 2:16, 3:14, 4:12}.get(lvl, 12)
            set_font(p.add_run(re.sub(r'\*\*(.+?)\*\*', r'\1', txt)), size=sz, bold=True, name_cn='黑体')
            i += 1; continue
        # 分隔线
        if re.match(r'^-{3,}$', ln): i += 1; continue
        # 引用
        if ln.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            add_bold_runs(p, ln.lstrip('> ').strip(), size=10.5, cn='楷体')
            for r_ in p.runs: r_.font.color.rgb = RGBColor(0x55,0x55,0x55)
            i += 1; continue
        # 列表
        m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.+)$', ln)
        if m:
            indent = len(m.group(1))
            p = doc.add_paragraph(style='List Bullet' if m.group(2) in '-*' else 'List Number')
            p.paragraph_format.left_indent = Cm(0.75 + indent * 0.4)
            p.paragraph_format.space_after = Pt(2)
            add_bold_runs(p, m.group(3), size=11)
            i += 1; continue
        # 空行
        if not ln.strip(): i += 1; continue
        # 正文
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_bold_runs(p, ln, size=12)
        i += 1

    doc.save(docx_path)
    return docx_path

if __name__ == '__main__':
    md = sys.argv[1]; out = sys.argv[2]
    title = sys.argv[3] if len(sys.argv) > 3 else '文档'
    subtitle = sys.argv[4] if len(sys.argv) > 4 else ''
    metas = sys.argv[5].split('|') if len(sys.argv) > 5 else []
    print(convert(md, out, title, subtitle, metas))
