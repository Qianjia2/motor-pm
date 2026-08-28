"""Excel export endpoints."""
import io
from datetime import date
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend_v2.database import get_db
from backend_v2.auth import get_current_user
from backend_v2.models import Project, WeeklyReport, RiskIssue

router = APIRouter(prefix="/api/export", tags=["export"])

HEADER_FILL = PatternFill(start_color="1d4ed8", end_color="1d4ed8", fill_type="solid")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def style_sheet(ws, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


@router.get("/projects")
def export_projects(db: Session = Depends(get_db), _current_user=Depends(get_current_user)):
    result = db.execute(
        select(Project).where(Project.is_active == True).order_by(Project.priority)
    )
    projects = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "项目总表"
    headers = ["编号", "名称", "客户", "优先级", "当前阶段", "完成率%", "状态", "难度", "开始日期", "计划结束"]
    style_sheet(ws, headers)

    for i, p in enumerate(projects, 2):
        ws.cell(row=i, column=1, value=p.code)
        ws.cell(row=i, column=2, value=p.name)
        ws.cell(row=i, column=3, value=p.client.name if p.client else "")
        ws.cell(row=i, column=4, value=p.priority)
        ws.cell(row=i, column=5, value=p.current_phase.name if p.current_phase else "")
        ws.cell(row=i, column=6, value=p.completion_pct)
        ws.cell(row=i, column=7, value=p.overall_status)
        ws.cell(row=i, column=8, value=p.difficulty)
        ws.cell(row=i, column=9, value=p.start_date.isoformat() if p.start_date else "")
        ws.cell(row=i, column=10, value=p.planned_end_date.isoformat() if p.planned_end_date else "")

    for col_letter, w in [("B", 30), ("C", 18)]:
        ws.column_dimensions[col_letter].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    fname = quote(f"项目总表_{date.today().isoformat()}.xlsx")
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fname}"})


@router.get("/projects/{project_id}/reports")
def export_reports(project_id: int, db: Session = Depends(get_db), _current_user=Depends(get_current_user)):
    result = db.execute(
        select(WeeklyReport).where(WeeklyReport.project_id == project_id)
        .order_by(WeeklyReport.year.desc(), WeeklyReport.week_number.desc())
    )
    reports = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    headers = ["年度", "周次", "报告日期", "总体进展", "关键成果", "技术线", "本周进展", "下周计划", "状态", "难点"]
    style_sheet(ws, headers)

    row = 2
    for r in reports:
        for li in r.line_items:
            line_name = li.line.short_name if li.line else f"Line{li.line_id}"
            ws.cell(row=row, column=1, value=r.year)
            ws.cell(row=row, column=2, value=r.week_number)
            ws.cell(row=row, column=3, value=r.report_date.isoformat() if r.report_date else "")
            ws.cell(row=row, column=4, value=r.overall_progress)
            ws.cell(row=row, column=5, value=r.key_accomplishments)
            ws.cell(row=row, column=6, value=line_name)
            ws.cell(row=row, column=7, value=li.this_week)
            ws.cell(row=row, column=8, value=li.next_week)
            ws.cell(row=row, column=9, value=li.status)
            ws.cell(row=row, column=10, value=li.blocker)
            row += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('周报_'+date.today().isoformat()+'.xlsx')}"})


@router.get("/projects/{project_id}/risks")
def export_risks(project_id: int, db: Session = Depends(get_db), _current_user=Depends(get_current_user)):
    result = db.execute(
        select(RiskIssue).where(RiskIssue.project_id == project_id).order_by(RiskIssue.created_at.desc())
    )
    risks = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    headers = ["类型", "标题", "描述", "概率", "影响", "等级", "状态", "缓解措施", "负责人", "识别日期", "目标解决", "验证结果"]
    style_sheet(ws, headers)

    for i, r in enumerate(risks, 2):
        ws.cell(row=i, column=1, value="风险" if r.type == "risk" else "问题")
        ws.cell(row=i, column=2, value=r.title)
        ws.cell(row=i, column=3, value=r.description or "")
        ws.cell(row=i, column=4, value=r.probability or "")
        ws.cell(row=i, column=5, value=r.impact or "")
        ws.cell(row=i, column=6, value=r.risk_level or "")
        ws.cell(row=i, column=7, value=r.status or "")
        ws.cell(row=i, column=8, value=r.mitigation_plan or "")
        ws.cell(row=i, column=9, value=r.owner.name if r.owner else "")
        ws.cell(row=i, column=10, value=r.identified_date.isoformat() if r.identified_date else "")
        ws.cell(row=i, column=11, value=r.target_resolve_date.isoformat() if r.target_resolve_date else "")
        ws.cell(row=i, column=12, value=r.resolution or "")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('风险清单_'+date.today().isoformat()+'.xlsx')}"})
