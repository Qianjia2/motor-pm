"""Project document CRUD + version lifecycle (P0/PP1/PP2)."""
from datetime import date, datetime
from fastapi import APIRouter, Depends, HTTPException, Query, Request, UploadFile, File, Form
from fastapi.responses import FileResponse
from sqlalchemy import select, update, delete
from sqlalchemy.orm import Session
from pathlib import Path
import mimetypes, aiofiles, uuid
from backend_v2.database import get_db
from backend_v2.auth import get_current_user
from backend_v2.audit import log_audit
from backend_v2.config import settings
from backend_v2.schemas import DocumentCreate, DocumentUpdate
from backend_v2.models import ProjectDocument, Project

router = APIRouter(tags=["documents"])

# Extensions whose file content should be stored as text for inline editing
_TEXT_EXTENSIONS = {
    "txt", "csv", "json", "xml", "yaml", "yml", "toml",
    "py", "ipynb", "js", "ts", "jsx", "tsx", "vue", "html", "htm", "css", "scss", "less",
    "c", "cpp", "h", "hpp", "cs", "java", "go", "rs", "swift", "kt",
    "sql", "sh", "bat", "ps1", "cmd", "bash", "zsh",
    "log", "cfg", "ini", "conf", "env", "properties", "editorconfig",
    "md", "rst", "markdown", "tex", "r", "rb", "php", "pl", "lua",
}


def doc_to_dict(d):
    r = {c.name: getattr(d, c.name) for c in d.__table__.columns}
    if r.get("upload_date"):
        r["upload_date"] = r["upload_date"].isoformat()
    return r


@router.get("/api/projects/{project_id}/docs")
def list_docs(
    project_id: int,
    doc_type: str = Query(None),
    status: str = Query(None),
    folder: str = Query(None),
    search: str = Query(""),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=10, le=200),
    db: Session = Depends(get_db),
    _current_user=Depends(get_current_user),
):
    q = select(ProjectDocument).where(ProjectDocument.project_id == project_id)
    if search:
        from sqlalchemy import or_
        q = q.where(or_(
            ProjectDocument.name.contains(search),
            ProjectDocument.notes.contains(search),
            ProjectDocument.folder.contains(search),
        ))
    if doc_type:
        q = q.where(ProjectDocument.doc_type == doc_type)
    if status:
        q = q.where(ProjectDocument.status == status)
    if folder is not None:
        if folder:
            from sqlalchemy import or_
            q = q.where(or_(
                ProjectDocument.folder == folder,
                ProjectDocument.folder.startswith(folder + "/"),
            ))
        else:
            q = q.where(ProjectDocument.folder == "")
    q = q.order_by(ProjectDocument.upload_date.desc())

    # Count total with filters
    from sqlalchemy import func
    count_filtered = select(func.count(ProjectDocument.id)).where(ProjectDocument.project_id == project_id)
    if search: count_filtered = count_filtered.where(or_(ProjectDocument.name.contains(search), ProjectDocument.notes.contains(search), ProjectDocument.folder.contains(search)))
    if doc_type: count_filtered = count_filtered.where(ProjectDocument.doc_type == doc_type)
    if status: count_filtered = count_filtered.where(ProjectDocument.status == status)
    if folder is not None:
        if folder:
            from sqlalchemy import or_
            count_filtered = count_filtered.where(or_(
                ProjectDocument.folder == folder,
                ProjectDocument.folder.startswith(folder + "/"),
            ))
        else:
            count_filtered = count_filtered.where(ProjectDocument.folder == "")
    total = db.execute(count_filtered).scalar() or 0

    offset = (page - 1) * page_size
    q = q.offset(offset).limit(page_size)
    result = db.execute(q)
    items = [doc_to_dict(d) for d in result.scalars().all()]
    # Mark broken references (path refs whose file no longer exists)
    for item in items:
        if item.get("notes", "").startswith("路径引用") and item.get("file_path"):
            import os as _os
            fp = item["file_path"]
            upload_root = Path(settings.UPLOAD_DIR).parent
            if not _os.path.exists(fp) and not _os.path.exists(str(upload_root / fp)):
                item["broken_ref"] = True

    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": max(1, (total + page_size - 1) // page_size),
    }


def _rebuild_kb_folder_cache():
    """Sync kb_folders.json with actual database folders after doc changes."""
    import json as _json, os as _os
    try:
        from backend_v2.database import SessionLocal
        from backend_v2.models import ProjectDocument, Project
        from sqlalchemy import select as _select
        db2 = SessionLocal()
        all_folders = set()
        docs = db2.execute(_select(ProjectDocument.folder, ProjectDocument.notes, ProjectDocument.project_id)).all()
        for folder_col, notes, proj_id in docs:
            f = folder_col
            if not f:
                try: f = _json.loads(notes or "{}").get("folder", "/")
                except: f = "/"
            if (not f or f == "/") and proj_id:
                proj = db2.get(Project, proj_id)
                if proj and proj.name:
                    f = "/" + proj.name
            if f and f != "/" and f != "":
                all_folders.add(f.replace("\\", "/"))
        cache_file = _os.path.join(str(Path(settings.UPLOAD_DIR).parent), "data", "kb_folders.json")
        paths = ["/"] + sorted(f for f in all_folders if f.startswith("/") or not f.startswith("/"))
        paths = ["/" if not p.startswith("/") else p for p in paths]
        with open(cache_file, "w", encoding="utf-8") as fp:
            _json.dump(paths, fp, ensure_ascii=False)
        db2.close()
    except Exception:
        pass  # Non-critical


def _load_folder_registry(project_id: int) -> list:
    """Load persisted folder list for a project (survives when docs deleted)."""
    import json, os
    fpath = os.path.join(settings.UPLOAD_DIR, "..", "data", "project_folders", f"{project_id}.json")
    try:
        with open(fpath, "r", encoding="utf-8") as f:
            return json.load(f)
    except:
        return []

def _save_folder_registry(project_id: int, folders: list):
    """Persist folder list for a project."""
    import json, os
    d = os.path.join(settings.UPLOAD_DIR, "..", "data", "project_folders")
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, f"{project_id}.json"), "w", encoding="utf-8") as f:
        json.dump(folders, f, ensure_ascii=False)


@router.get("/api/projects/{project_id}/doc-folders")
def list_doc_folders(
    project_id: int,
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
):
    """List folder tree with counts. Empty folders from registry also appear."""
    from sqlalchemy import func
    result = db.execute(
        select(ProjectDocument.folder, func.count(ProjectDocument.id))
        .where(ProjectDocument.project_id == project_id)
        .group_by(ProjectDocument.folder)
    )
    rows = [(r[0] or "", r[1]) for r in result.all()]

    # Merge persisted registry (empty folders survive doc deletion)
    registered = set(_load_folder_registry(project_id))
    for f in registered:
        if not any(r[0] == f for r in rows):
            rows.append((f, 0))

    tree = {}
    for folder, count in rows:
        if folder:
            parts = folder.split("/")
            prefix = ""
            for i, p in enumerate(parts):
                full = prefix + "/" + p if prefix else p
                if full not in tree:
                    tree[full] = {"name": p, "level": i, "count": 0, "path": full}
                if i == len(parts) - 1:
                    tree[full]["count"] += count
                prefix = full
    result_tree = []
    for path, info in sorted(tree.items()):
        parent = path.rsplit("/", 1)[0] if "/" in path else ""
        if parent and parent in tree:
            tree[parent]["count"] += info["count"]
        result_tree.append(info)
    return {
        "folders": sorted(set(f for f, _ in rows if f)),
        "tree": result_tree,
        "root_count": sum(c for f, c in rows if not f) or 0,
    }


@router.post("/api/projects/{project_id}/doc-folders")
def save_doc_folder(project_id: int, data: dict, _user=Depends(get_current_user)):
    """Register a folder so it persists even when empty."""
    folders = _load_folder_registry(project_id)
    name = (data.get("name") or "").strip()
    if name and name not in folders:
        folders.append(name)
        _save_folder_registry(project_id, folders)
    return {"ok": True, "folders": folders}


@router.post("/api/projects/{project_id}/docs", status_code=201)
async def create_doc(
    project_id: int,
    name: str = Form(...),
    doc_type: str = Form(...),
    folder: str = Form(""),
    notes: str = Form(""),
    file: UploadFile = File(None),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    file_path = None
    file_size = 0
    text_content = None
    if file and file.filename:
        ext = Path(file.filename).suffix.lstrip(".").lower()
        if ext not in settings.ALLOWED_EXTENSIONS:
            raise HTTPException(status_code=400, detail=f"不支持的文件类型: .{ext}")
        upload_dir = Path(settings.UPLOAD_DIR) / str(project_id) / "docs"
        upload_dir.mkdir(parents=True, exist_ok=True)
        safe_name = f"{uuid.uuid4().hex[:8]}_{Path(file.filename).name}"
        dest = upload_dir / safe_name
        raw_bytes = await file.read()
        if len(raw_bytes) > settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024:
            raise HTTPException(status_code=413, detail=f"文件超过 {settings.MAX_UPLOAD_SIZE_MB}MB 限制")
        with open(dest, "wb") as f:
            f.write(raw_bytes)
        file_path = str(dest.relative_to(Path(settings.UPLOAD_DIR).parent))
        file_size = len(raw_bytes)
        # Extract text content for text-based files so they can be edited inline
        if ext in _TEXT_EXTENSIONS:
            try:
                text_content = raw_bytes.decode('utf-8')
            except UnicodeDecodeError:
                pass  # not valid UTF-8, leave as file-only

    d = ProjectDocument(
        project_id=project_id, name=name, doc_type=doc_type,
        folder=folder,
        file_path=file_path, file_size=file_size,
        content=text_content,
        uploaded_by=current_user.username, upload_date=date.today(),
        notes=notes, version=1,
    )
    db.add(d)
    db.commit()
    db.refresh(d)
    log_audit(db, current_user.username, "create", "document", d.id, d.name, project_id, f"上传文档: {d.name}")
    db.commit()
    _rebuild_kb_folder_cache()
    return doc_to_dict(d)


@router.put("/api/docs/{doc_id}")
def update_doc(
    doc_id: int, data: DocumentUpdate,
    db: Session = Depends(get_db), current_user=Depends(get_current_user),
):
    result = db.execute(select(ProjectDocument).where(ProjectDocument.id == doc_id))
    d = result.scalar_one_or_none()
    if not d:
        raise HTTPException(status_code=404, detail="不存在")
    upd = data.model_dump(exclude_unset=True)
    db.execute(update(ProjectDocument).where(ProjectDocument.id == doc_id).values(**upd))
    db.commit()
    return {"message": "已更新"}


@router.delete("/api/docs/{doc_id}")
def delete_doc(
    doc_id: int, db: Session = Depends(get_db), current_user=Depends(get_current_user),
):
    result = db.execute(select(ProjectDocument).where(ProjectDocument.id == doc_id))
    d = result.scalar_one_or_none()
    if not d:
        raise HTTPException(status_code=404, detail="不存在")
    db.execute(delete(ProjectDocument).where(ProjectDocument.id == doc_id))
    log_audit(db, current_user.username, "delete", "document", doc_id, d.name, d.project_id, f"删除文档: {d.name}")
    db.commit()
    return {"message": "已删除"}


@router.post("/api/docs/{doc_id}/deprecate")
def deprecate_doc(
    doc_id: int, db: Session = Depends(get_db), current_user=Depends(get_current_user),
):
    """Mark a document as deprecated."""
    result = db.execute(select(ProjectDocument).where(ProjectDocument.id == doc_id))
    d = result.scalar_one_or_none()
    if not d:
        raise HTTPException(status_code=404, detail="不存在")
    db.execute(update(ProjectDocument).where(ProjectDocument.id == doc_id).values(status="deprecated"))
    log_audit(db, current_user.username, "deprecate", "document", doc_id, d.name, d.project_id, f"标记失效: {d.name}")
    db.commit()
    return {"message": "已标记为失效"}


@router.post("/api/docs/{doc_id}/replace", status_code=201)
async def replace_doc(
    doc_id: int,
    file: UploadFile = File(...),
    notes: str = Form(""),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Upload a new version — old one auto-superseded."""
    result = db.execute(select(ProjectDocument).where(ProjectDocument.id == doc_id))
    old = result.scalar_one_or_none()
    if not old:
        raise HTTPException(status_code=404, detail="不存在")

    ext = Path(file.filename).suffix.lstrip(".").lower()
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: .{ext}")

    upload_dir = Path(settings.UPLOAD_DIR) / str(old.project_id) / "docs"
    upload_dir.mkdir(parents=True, exist_ok=True)
    safe_name = f"{uuid.uuid4().hex[:8]}_{Path(file.filename).name}"
    dest = upload_dir / safe_name
    raw_bytes = await file.read()
    if len(raw_bytes) > settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"文件超过 {settings.MAX_UPLOAD_SIZE_MB}MB 限制")
    with open(dest, "wb") as f:
        f.write(raw_bytes)

    text_content = None
    if ext in _TEXT_EXTENSIONS:
        try:
            text_content = raw_bytes.decode('utf-8')
        except UnicodeDecodeError:
            pass

    new_ver = old.version + 1
    new_doc = ProjectDocument(
        project_id=old.project_id, name=old.name, doc_type=old.doc_type,
        file_path=str(dest.relative_to(Path(settings.UPLOAD_DIR).parent)),
        file_size=len(raw_bytes),
        content=text_content,
        uploaded_by=current_user.username, upload_date=date.today(),
        notes=notes, version=new_ver, folder=getattr(old, 'folder', '/'),
    )
    db.add(new_doc)
    db.flush()

    # Mark old as superseded
    db.execute(update(ProjectDocument).where(ProjectDocument.id == doc_id).values(
        status="superseded", replaced_by_id=new_doc.id,
    ))

    log_audit(db, current_user.username, "replace", "document", new_doc.id, f"V{new_ver}: {old.name}", old.project_id, f"替换文档: V{old.version}→V{new_ver}")
    db.commit()
    db.refresh(new_doc)
    return doc_to_dict(new_doc)


@router.get("/api/docs/{doc_id}/preview")
def preview_doc(
    doc_id: int,
    raw: bool = Query(False),
    token: str = Query(""),
    db: Session = Depends(get_db),
    request: Request = None,
):
    """Preview a document: converts to HTML for inline preview, or serves raw file."""
    from fastapi.responses import HTMLResponse

    # Auth: check token param (for iframe embed) or Authorization header
    jwt_token = token or ""
    if not jwt_token and request:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            jwt_token = auth[7:]

    # Only enforce auth for non-raw access; raw is used inside iframe on already-authed page
    if not raw:
        if not jwt_token:
            raise HTTPException(status_code=401, detail="需要登录")
        from backend_v2.auth import verify_token
        try:
            verify_token(jwt_token)
        except Exception:
            raise HTTPException(status_code=401, detail="需要登录")

    result = db.execute(select(ProjectDocument).where(ProjectDocument.id == doc_id))
    d = result.scalar_one_or_none()
    if not d:
        raise HTTPException(status_code=404, detail="文档不存在")

    # Resolve file path
    fpath = d.file_path
    upload_root = Path(settings.UPLOAD_DIR).parent
    if fpath:
        full = upload_root / fpath
        if full.exists():
            fpath = str(full)
        elif not Path(fpath).exists():
            fpath = None

    if not fpath:
        text = (d.content or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        html = f"<html><head><meta charset='utf-8'><style>body{{font-family:sans-serif;padding:24px;line-height:1.8;white-space:pre-wrap;word-break:break-word;font-size:14px}}</style></head><body><h3>{d.name}</h3>{text}</body></html>"
        return HTMLResponse(html)

    ext = fpath.rsplit(".", 1)[-1].lower() if "." in fpath else ""

    # Images: serve inline
    if ext in ("png", "jpg", "jpeg", "gif", "webp", "bmp", "svg"):
        mime = mimetypes.guess_type(fpath)[0] or "image/png"
        return FileResponse(fpath, media_type=mime)

    # Raw file — serve inline for preview (not download)
    if raw:
        mime, _ = mimetypes.guess_type(fpath)
        return FileResponse(fpath, media_type=mime or "application/octet-stream",
                           headers={"Content-Disposition": "inline"})

    html = ""
    # Word (.docx)
    if ext in ("docx",):
        try:
            from docx import Document as DocxDoc
            import zipfile, base64 as _b64, io as _io2
            doc = DocxDoc(fpath)
            parts = []
            image_map = {}
            try:
                with zipfile.ZipFile(fpath) as zf:
                    for n in zf.namelist():
                        if "media" in n.lower() and any(n.lower().endswith(e) for e in (".png",".jpg",".jpeg",".gif",".bmp",".webp")):
                            b64 = _b64.b64encode(zf.read(n)).decode()
                            ext_i = n.rsplit(".",1)[-1].lower()
                            m = {"png":"image/png","jpg":"image/jpeg","jpeg":"image/jpeg","gif":"image/gif","webp":"image/webp","bmp":"image/bmp"}.get(ext_i,"image/png")
                            image_map[n] = f"data:{m};base64,{b64}"
            except: pass

            for p in doc.paragraphs:
                text = ""
                for run in p.runs:
                    if run.text:
                        t = run.text.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
                        if run.bold: t = f"<b>{t}</b>"
                        if run.italic: t = f"<i>{t}</i>"
                        text += t
                if p.style.name.startswith("Heading"):
                    level = p.style.name.replace("Heading ","").strip()
                    parts.append(f"<h{level}>{text}</h{level}>")
                else:
                    parts.append(f"<p>{text or '&nbsp;'}</p>")
            for table in doc.tables:
                rows = []
                for row in table.rows:
                    cells = "".join(f"<td style='border:1px solid #ddd;padding:6px 10px'>{c.text.replace('&','&amp;').replace('<','&lt;')}</td>" for c in row.cells)
                    rows.append(f"<tr>{cells}</tr>")
                parts.append(f"<table style='border-collapse:collapse;width:100%;margin:8px 0'>{''.join(rows)}</table>")
            html = "".join(parts)
        except Exception as e:
            html = f"<p style='color:#e6a23c'>Word 解析失败: {e}</p>"

    # PDF — embed directly for full fidelity
    elif ext == "pdf":
        raw_url = f"/api/docs/{doc_id}/preview?raw=1"
        html = f"<iframe src='{raw_url}' style='width:100%;height:70vh;border:none;border-radius:4px'></iframe>"

    # Excel
    elif ext in ("xlsx", "xls", "csv"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(fpath, data_only=True)
            all_parts = []
            for sn in wb.sheetnames[:5]:
                ws = wb[sn]
                rows_html = []
                for row in ws.iter_rows(max_row=100, values_only=True):
                    cells = "".join(f"<td style='border:1px solid #ddd;padding:4px 8px;font-size:12px'>{str(c or '').replace('&','&amp;').replace('<','&lt;')}</td>" for c in row)
                    rows_html.append(f"<tr>{cells}</tr>")
                all_parts.append(f"<h4 style='margin:12px 0 4px'>📋 {sn}</h4><table style='border-collapse:collapse;width:100%;margin-bottom:16px'>{''.join(rows_html)}</table>")
            html = "".join(all_parts)
        except Exception as e:
            html = f"<p style='color:#e6a23c'>Excel 解析失败: {e}</p>"

    # Text / code
    elif ext in ("txt", "md", "py", "js", "ts", "vue", "json", "xml", "yaml", "yml", "html", "css", "sql", "sh", "bat"):
        try:
            for enc in ("utf-8", "gbk", "gb18030", "latin-1"):
                try:
                    with open(fpath, "r", encoding=enc) as f:
                        text = f.read()
                    break
                except: continue
            text = text.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
            html = f"<pre style='white-space:pre-wrap;font-family:Consolas,monospace;font-size:12px;line-height:1.6;background:#f5f7fa;padding:16px;border-radius:4px;overflow:auto;max-height:65vh'>{text}</pre>"
        except Exception as e:
            html = f"<p style='color:#e6a23c'>文件读取失败: {e}</p>"

    else:
        # Unsupported type — offer download
        html = f"<div style='text-align:center;padding:60px 20px'><p style='font-size:48px;margin:0'>📁</p><p style='font-size:16px;color:#606266;margin:16px 0'>{d.name}</p><p style='color:#909399'>此文件类型 (.{ext}) 不支持在线预览</p><a href='/api/docs/{doc_id}/preview?raw=1' style='display:inline-block;margin-top:12px;padding:8px 24px;background:#409eff;color:#fff;border-radius:4px;text-decoration:none'>⬇ 下载文件</a></div>"

    full_html = f"<html><head><meta charset='utf-8'><style>body{{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;padding:20px 24px;line-height:1.8;max-width:960px;margin:0 auto;font-size:14px;color:#303133}} h1,h2,h3{{margin-top:20px}} table{{font-size:13px}} img{{max-width:100%}}</style></head><body>{html}</body></html>"
    return HTMLResponse(full_html)


@router.post("/api/projects/{project_id}/docs/batch")
def batch_operate_docs(project_id: int, data: dict, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    """Batch delete or move project documents."""
    ids = data.get("ids", [])
    action = data.get("action", "")
    if not ids or action not in ("delete", "move"):
        raise HTTPException(400, "参数错误")
    if action == "delete":
        db.execute(delete(ProjectDocument).where(ProjectDocument.id.in_(ids), ProjectDocument.project_id == project_id))
        db.commit()
        return {"message": f"deleted {len(ids)}"}
    if action == "move":
        folder = data.get("folder", "")
        values = {"folder": folder}
        db.execute(update(ProjectDocument).where(ProjectDocument.id.in_(ids), ProjectDocument.project_id == project_id).values(**values))
        db.commit()
        return {"message": f"moved {len(ids)}"}


@router.post("/api/projects/{project_id}/docs/add-ref", status_code=201)
def add_doc_ref(
    project_id: int,
    data: dict,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Add a document by reference (file path only, no copy)."""
    file_path = data.get("file_path", "")
    name = data.get("name", "") or Path(file_path).stem
    doc_type = data.get("doc_type", "pp2_report")
    folder = data.get("folder", "")
    notes = data.get("notes", f"路径引用: {file_path}")

    if not file_path:
        raise HTTPException(400, "缺少文件路径")

    d = ProjectDocument(
        project_id=project_id, name=name, doc_type=doc_type,
        folder=folder,
        file_path=file_path.replace("\\", "/"),
        file_size=0,
        uploaded_by=current_user.username, upload_date=date.today(),
        notes=notes, version=1,
    )
    db.add(d)
    db.commit()
    db.refresh(d)
    _rebuild_kb_folder_cache()
    return doc_to_dict(d)

# ── Server file browser & import ──

@router.get("/api/docs/browse-server")
def browse_server_dir(
    path: str = Query(""),
    _user=Depends(get_current_user),
):
    """Browse server/inner-network directories for file import."""
    import os as _os
    data_root = str(Path(settings.UPLOAD_DIR).parent)
    app_root = str(Path(__file__).parent.parent.parent)

    current = _os.path.normpath(path) if path else data_root
    # Allow drive letters (C:\) and UNC paths
    if not _os.path.exists(current):
        current = data_root

    dirs, files = [], []
    try:
        for name in sorted(_os.listdir(current)):
            if name.startswith("."): continue
            full = _os.path.join(current, name)
            try:
                if _os.path.isdir(full):
                    dirs.append({"name": name, "path": full.replace("\\", "/")})
                elif _os.path.isfile(full):
                    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
                    files.append({"name": name, "path": full.replace("\\", "/"), "ext": ext, "size": _os.path.getsize(full)})
            except (PermissionError, OSError):
                pass
    except PermissionError:
        raise HTTPException(403, "无权访问该目录")

    parent = _os.path.dirname(current).replace("\\", "/")
    if parent == current.replace("\\", "/"):
        parent = None  # Root of a drive
    return {
        "current": current.replace("\\", "/"),
        "parent": parent,
        "roots": [data_root.replace("\\", "/"), app_root.replace("\\", "/")],
        "dirs": dirs,
        "files": files,
    }


@router.post("/api/projects/{project_id}/docs/import-server", status_code=201)
def import_server_file(
    project_id: int,
    file_path: str = Form(...),
    doc_type: str = Form(""),
    name: str = Form(""),
    notes: str = Form(""),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Import a file from server/local path into project documents."""
    import shutil
    src = Path(file_path)
    if not src.exists() or not src.is_file():
        raise HTTPException(404, f"文件不存在: {file_path}")
    if not name:
        name = src.stem
    ext = src.suffix.lstrip(".").lower()
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise HTTPException(400, f"不支持的文件类型: .{ext}")

    project_check = db.execute(select(Project).where(Project.id == project_id)).scalar_one_or_none()
    if not project_check:
        raise HTTPException(404, "项目不存在")
    if not doc_type:
        doc_type = "pp2_report"

    dest_dir = Path(settings.UPLOAD_DIR) / str(project_id) / "docs"
    dest_dir.mkdir(parents=True, exist_ok=True)
    safe_name = f"{uuid.uuid4().hex[:8]}_{src.name}"
    dest = dest_dir / safe_name
    shutil.copy2(src, dest)

    d = ProjectDocument(
        project_id=project_id, name=name, doc_type=doc_type,
        file_path=str(dest.relative_to(Path(settings.UPLOAD_DIR).parent)),
        file_size=dest.stat().st_size,
        uploaded_by=current_user.username, upload_date=date.today(),
        notes=notes or f"从服务器导入: {file_path}", version=1, folder="/",
    )
    db.add(d)
    db.commit()
    db.refresh(d)
    log_audit(db, current_user.username, "import_server", "document", d.id, d.name, project_id, f"从服务器导入: {d.name}")
    db.commit()
    return doc_to_dict(d)
