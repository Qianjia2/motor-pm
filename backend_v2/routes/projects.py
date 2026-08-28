"""Project CRUD + copy + stats."""
from datetime import date, datetime, timedelta
from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import select, update, delete, func, or_, case
from sqlalchemy.orm import Session, selectinload
from backend_v2.database import get_db
from backend_v2.auth import get_current_user
from backend_v2.audit import log_audit
from backend_v2.schemas import (
    ProjectCreate, ProjectUpdate, CopyProjectRequest, PaginatedResponse,
)
from backend_v2.models import Project, ProjectPhaseGate, Phase, Gate, Milestone, Deliverable, RiskIssue, GateDeliverableStandard

router = APIRouter(prefix="/api/projects", tags=["projects"])


def derive_project_health(db: Session, projects: list) -> dict:
    """按健康信号实时推导项目整体状态(不落库):
    - 已完成项目(有实际完成日期) → completed
    - 逾期里程碑>=3 或 未关闭风险>=2 → blocked
    - 逾期里程碑>=1 或 未关闭风险>=1 → at_risk
    - 其余 → normal
    """
    pids = [p.id for p in projects if p.id]
    if not pids:
        return {}
    today = date.today()
    overdue_map = dict(db.execute(
        select(Milestone.project_id, func.count(Milestone.id)).where(
            Milestone.project_id.in_(pids),
            Milestone.status != "completed",
            Milestone.planned_date < today,
        ).group_by(Milestone.project_id)
    ).all())
    # pending_review 为 AI 待审核,同样属于未关闭
    open_risk_map = dict(db.execute(
        select(RiskIssue.project_id, func.count(RiskIssue.id)).where(
            RiskIssue.project_id.in_(pids),
            RiskIssue.status.in_(["open", "in_progress", "pending_review"]),
        ).group_by(RiskIssue.project_id)
    ).all())
    health = {}
    for p in projects:
        if p.actual_end_date:
            health[p.id] = "completed"
            continue
        overdue = overdue_map.get(p.id, 0)
        risks = open_risk_map.get(p.id, 0)
        if overdue >= 3 or risks >= 2:
            health[p.id] = "blocked"
        elif overdue >= 1 or risks >= 1:
            health[p.id] = "at_risk"
        else:
            health[p.id] = "normal"
    return health


def _fix_empty_strings(data: dict) -> dict:
    """Replace empty strings with None for nullable fields."""
    for k, v in data.items():
        if v == "":
            data[k] = None
    return data


def project_to_dict(p: Project, milestone_pct: int = None) -> dict:
    d = {c.name: getattr(p, c.name) for c in p.__table__.columns}
    for key in ("start_date", "planned_end_date", "actual_end_date", "created_at", "updated_at"):
        val = d.get(key)
        d[key] = val.isoformat() if val else None
    d["current_phase"] = {"id": p.current_phase.id, "name": p.current_phase.name} if p.current_phase else None
    d["client"] = {"id": p.client.id, "name": p.client.name, "abbreviation": p.client.abbreviation} if p.client else None
    d["member_count"] = len(p.members) if p.members else 0
    pm_name = None
    for m in p.members or []:
        if m.role and m.role.name == "项目经理":
            pm_name = m.member.name
            break
    d["pm"] = pm_name
    if milestone_pct is not None:
        d["completion_pct"] = milestone_pct
    return d


@router.get("", response_model=PaginatedResponse)
def list_projects(
    status: str = Query(None, description="filter by overall_status"),
    priority: str = Query(None, description="filter by priority"),
    client_id: int = Query(None),
    search: str = Query(None, description="full-text search"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    sort_by: str = Query("priority", description="sort field"),
    db: Session = Depends(get_db),
    _current_user=Depends(get_current_user),
):
    q = select(Project).where(Project.is_active == True).options(
        selectinload(Project.client),
        selectinload(Project.current_phase),
        selectinload(Project.members),
    )

    # Non-admin users: only see projects they are assigned to
    if _current_user.role != "admin":
        from backend_v2.models import ProjectMember
        assigned_project_ids = db.execute(
            select(ProjectMember.project_id).where(ProjectMember.member_id == _current_user.member_id)
        ).scalars().all()
        if assigned_project_ids:
            q = q.where(Project.id.in_(assigned_project_ids))
        else:
            q = q.where(Project.id == -1)  # no projects assigned

    if status:
        q = q.where(Project.overall_status == status)
    if priority:
        q = q.where(Project.priority == priority)
    if client_id:
        q = q.where(Project.client_id == client_id)
    if search:
        q = q.where(
            (Project.name.ilike(f"%{search}%")) | (Project.code.ilike(f"%{search}%"))
        )

    # Count
    count_q = select(func.count()).select_from(q.subquery())
    total = (db.execute(count_q)).scalar()

    # Sort（白名单：避免 getattr 取到关系属性导致 500）
    ALLOWED_SORT_COLS = {"id", "name", "code", "priority", "overall_status",
                         "start_date", "planned_end_date", "completion_pct",
                         "created_at", "updated_at"}
    if sort_by not in ALLOWED_SORT_COLS:
        sort_by = "priority"
    sort_col = getattr(Project, sort_by)
    q = q.order_by(sort_col).offset((page - 1) * page_size).limit(page_size)

    result = db.execute(q)
    projects = result.scalars().all()

    # Calculate completion_pct + current_phase from milestones
    pct_map = {}
    phase_map = {}  # project_id -> current_phase_name
    if projects:
        pids = [p.id for p in projects]
        from backend_v2.models import Milestone, Phase
        ms_query = select(
            Milestone.project_id,
            func.count(Milestone.id).label("total"),
            func.sum(case((Milestone.status == "completed", 1), else_=0)).label("done"),
        ).where(Milestone.project_id.in_(pids)).group_by(Milestone.project_id)
        for row in db.execute(ms_query):
            total_ms = row.total or 0; done_ms = row.done or 0
            pct_map[row.project_id] = round(done_ms / total_ms * 100) if total_ms > 0 else 0

        # Auto-detect current phase: first phase (by sort_order) with incomplete milestones
        phases_all = db.execute(select(Phase).order_by(Phase.sort_order)).scalars().all()
        phase_order = {p.id: (p.name, i) for i, p in enumerate(phases_all)}
        ms_phase_query = select(
            Milestone.project_id,
            Milestone.phase_id,
            func.count(Milestone.id).label("total"),
            func.sum(case((Milestone.status == "completed", 1), else_=0)).label("done"),
        ).where(Milestone.project_id.in_(pids)).group_by(Milestone.project_id, Milestone.phase_id)
        proj_phase_ms = {}
        for row in db.execute(ms_phase_query):
            proj_phase_ms.setdefault(row.project_id, []).append({
                "phase_id": row.phase_id, "total": row.total, "done": row.done,
            })
        for pid in pids:
            items = proj_phase_ms.get(pid, [])
            if not items:
                continue
            # Sort by phase order
            items.sort(key=lambda x: phase_order.get(x["phase_id"], ("?", 999))[1])
            # Find first phase with incomplete milestones
            current_name = None
            for it in items:
                if it["done"] < it["total"]:
                    current_name = phase_order.get(it["phase_id"], (None, 999))[0]
                    break
            if not current_name:
                current_name = phase_order.get(items[-1]["phase_id"], (None,))[0]  # last phase
            phase_map[pid] = current_name

    health_map = derive_project_health(db, projects)

    def _project_dict(p):
        d = project_to_dict(p, pct_map.get(p.id))
        # 推导后的健康状态覆盖落库字段(全库默认 normal,无维护入口)
        d["overall_status"] = health_map.get(p.id, d["overall_status"])
        # Override current_phase with auto-detected one
        auto_phase = phase_map.get(p.id)
        if auto_phase:
            d["current_phase"] = {"id": None, "name": auto_phase}
        elif not d.get("current_phase"):
            d["current_phase"] = {"id": None, "name": "未开始"}
        return d

    return PaginatedResponse(
        data=[_project_dict(p) for p in projects],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/{project_id}/plan")
def generate_project_plan(project_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    """Generate a project plan document from existing project data (Excel format)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from io import BytesIO
    from fastapi.responses import StreamingResponse
    from backend_v2.models import Milestone as MS, RiskIssue as RI, Requirement as RQ, Task as TK

    p_result = db.execute(
        select(Project).where(Project.id == project_id).options(
            selectinload(Project.client), selectinload(Project.current_phase),
        )
    )
    p = p_result.scalars().one_or_none()
    if not p: raise HTTPException(status_code=404, detail="项目不存在")

    # Load related data separately to avoid lazy-load issues
    ms_result = db.execute(select(MS).where(MS.project_id == project_id).order_by(MS.phase_id, MS.planned_date))
    milestones = ms_result.scalars().all()
    risk_result = db.execute(select(RI).where(RI.project_id == project_id).order_by(RI.severity.desc()))
    risks = risk_result.scalars().all()
    req_result = db.execute(select(RQ).where(RQ.project_id == project_id).order_by(RQ.priority, RQ.created_at))
    requirements = req_result.scalars().all()
    task_result = db.execute(select(TK).where(TK.project_id == project_id).order_by(TK.sort_order))
    tasks = task_result.scalars().all()

    wb = Workbook()
    title_font = Font(name='微软雅黑', size=16, bold=True)
    h2_font = Font(name='微软雅黑', size=12, bold=True)
    cell_font = Font(name='微软雅黑', size=11)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    header_fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')

    def write_info(ws, row, label, value):
        ws.cell(row=row, column=1, value=label).font = Font(name='微软雅黑', size=11, bold=True)
        ws.cell(row=row, column=2, value=str(value or '')).font = cell_font
        return row + 1

    # Sheet 1: Project Overview
    ws1 = wb.active; ws1.title = "项目基本信息"
    ws1.merge_cells('A1:D1'); ws1.cell(row=1, column=1, value=f"项目计划书：{p.name}").font = title_font
    ws1.merge_cells('A2:D2'); ws1.cell(row=2, column=1, value=f"编号：{p.code} | 生成日期：{date.today().isoformat()}").font = Font(name='微软雅黑', size=10, color='6B7280')
    r = 4
    r = write_info(ws1, r, "项目名称", p.name)
    r = write_info(ws1, r, "项目编号", p.code)
    r = write_info(ws1, r, "客户", p.client.name if p.client else '-')
    r = write_info(ws1, r, "当前阶段", p.current_phase.name if p.current_phase else '-')
    r = write_info(ws1, r, "优先级", p.priority or '-')
    r = write_info(ws1, r, "计划开始", p.start_date.isoformat() if p.start_date else '-')
    r = write_info(ws1, r, "计划完成", p.planned_end_date.isoformat() if p.planned_end_date else '-')
    r = write_info(ws1, r + 1, "项目描述", p.description or '-')
    r = write_info(ws1, r, "最终交付物", p.final_deliverable or '-')
    r = write_info(ws1, r, "验收标准", p.acceptance_criteria or '-')
    r = write_info(ws1, r + 1, "合同金额", f"{p.contract_amount or 0}万元")
    r = write_info(ws1, r, "项目预算", f"{p.budget or 0}万元")
    ws1.column_dimensions['A'].width = 15; ws1.column_dimensions['B'].width = 50

    # Sheet 2: Milestones
    ws2 = wb.create_sheet("里程碑计划")
    ws2.append(["阶段", "名称", "计划日期", "结束日期", "状态", "关键节点"])
    for cell in ws2[1]: cell.font = header_font; cell.fill = header_fill
    for ms in milestones:
        ws2.append([ms.phase.name if ms.phase else '-', ms.name, ms.planned_date.isoformat() if ms.planned_date else '-',
                     ms.planned_end_date.isoformat() if ms.planned_end_date else '-', ms.status or '-', '是' if ms.is_key else ''])
    ws2.column_dimensions['A'].width = 14; ws2.column_dimensions['B'].width = 30

    # Sheet 3: Requirements
    ws3 = wb.create_sheet("需求清单")
    ws3.append(["标题", "分类", "优先级", "来源", "状态"])
    for cell in ws3[1]: cell.font = header_font; cell.fill = header_fill
    for rq in requirements:
        ws3.append([rq.title, rq.category or '-', rq.priority or '-', rq.source or '-', rq.status or '-'])
    ws3.column_dimensions['A'].width = 40

    # Sheet 4: WBS Tasks
    ws4 = wb.create_sheet("WBS任务分解")
    ws4.append(["层级", "任务名称", "负责人", "进度%", "状态", "截止日期"])
    for cell in ws4[1]: cell.font = header_font; cell.fill = header_fill
    for tk in tasks:
        ws4.append([(tk.parent_id and 'L1' or 'L0'), tk.name, tk.assignee.name if tk.assignee else '-',
                     tk.progress_pct or 0, tk.status or '-', tk.end_date.isoformat() if tk.end_date else '-'])
    ws4.column_dimensions['B'].width = 35

    # Sheet 5: Risks
    ws5 = wb.create_sheet("风险登记")
    ws5.append(["标题", "严重度", "概率", "影响", "缓解措施", "状态"])
    for cell in ws5[1]: cell.font = header_font; cell.fill = header_fill
    for ri in sorted(risks, key=lambda x: -(int(x.severity) if x.severity and x.severity.isdigit() else 0)):
        ws5.append([ri.title, ri.severity or '-', ri.probability or '-', ri.impact or '-', ri.mitigation_plan or '-', ri.status or '-'])
    ws5.column_dimensions['A'].width = 35; ws5.column_dimensions['E'].width = 40

    output = BytesIO(); wb.save(output); output.seek(0)
    safe_name = p.code.replace('/', '-').replace('\\', '-') if p.code else f"Project-{project_id}"
    return StreamingResponse(output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=ProjectPlan-{safe_name}.xlsx"})


@router.get("/search")
def global_search(
    q: str = Query(..., min_length=1, max_length=100),
    db: Session = Depends(get_db), _current_user=Depends(get_current_user),
):
    """Global search across projects, milestones, documents, and risks."""
    from backend_v2.models import Milestone as MS, RiskIssue as RI, ProjectDocument as PD
    like = f"%{q}%"
    results = {"projects": [], "milestones": [], "documents": [], "risks": []}

    proj_result = db.execute(
        select(Project).options(selectinload(Project.current_phase))
        .where(Project.is_active == True, or_(Project.name.like(like), Project.code.like(like))).limit(8)
    )
    for p in proj_result.scalars().all():
        results["projects"].append({"id": p.id, "name": p.name, "code": p.code, "phase": p.current_phase.name if p.current_phase else None})

    ms_result = db.execute(
        select(MS).options(selectinload(MS.project))
        .join(Project).where(Project.is_active == True, MS.name.like(like)).limit(8)
    )
    for m in ms_result.scalars().all():
        results["milestones"].append({"id": m.id, "name": m.name, "project_id": m.project_id, "project_name": m.project.name if m.project else None, "planned_date": m.planned_date.isoformat() if m.planned_date else None})

    doc_result = db.execute(
        select(PD).options(selectinload(PD.project))
        .join(Project).where(Project.is_active == True, PD.name.like(like)).limit(8)
    )
    for d in doc_result.scalars().all():
        results["documents"].append({"id": d.id, "name": d.name, "project_id": d.project_id, "project_name": d.project.name if d.project else None, "doc_type": d.doc_type})

    risk_result = db.execute(
        select(RI).options(selectinload(RI.project))
        .join(Project).where(Project.is_active == True, or_(RI.title.like(like), RI.description.like(like))).limit(8)
    )
    for r in risk_result.scalars().all():
        results["risks"].append({"id": r.id, "title": r.title, "project_id": r.project_id, "project_name": r.project.name if r.project else None, "severity": r.severity, "status": r.status})

    return results


@router.get("/{project_id}")
def get_project(project_id: int, db: Session = Depends(get_db), _current_user=Depends(get_current_user)):
    result = db.execute(
        select(Project).where(Project.id == project_id, Project.is_active == True).options(
            selectinload(Project.client),
            selectinload(Project.current_phase),
            selectinload(Project.members),
            selectinload(Project.phase_gates),
            selectinload(Project.milestones),
        )
    )
    p = result.scalars().one_or_none()
    if not p:
        raise HTTPException(status_code=404, detail="项目不存在")

    # Non-admin users: only access assigned projects
    if _current_user.role != "admin" and _current_user.member_id:
        member_ids = [pm.member_id for pm in (p.members or [])]
        if _current_user.member_id not in member_ids:
            raise HTTPException(status_code=403, detail="无权访问此项目")
    ms_list = p.milestones or []
    total_ms = len(ms_list)
    done_ms = sum(1 for m in ms_list if m.status == "completed")
    pct = round(done_ms / total_ms * 100) if total_ms > 0 else 0
    d = project_to_dict(p, pct)
    d["overall_status"] = derive_project_health(db, [p]).get(p.id, d["overall_status"])
    d["phase_gates"] = [_pg_to_dict(pg) for pg in (p.phase_gates or [])]
    return d


def _pg_to_dict(pg):
    d = {c.name: getattr(pg, c.name) for c in pg.__table__.columns}
    for df in ("planned_start_date", "planned_end_date", "actual_start_date", "actual_end_date", "gate_review_date"):
        d[df] = d[df].isoformat() if d.get(df) else None
    d["phase"] = {"id": pg.phase.id, "name": pg.phase.name, "code": pg.phase.code} if pg.phase else None
    d["gate"] = {"id": pg.gate.id, "name": pg.gate.name, "code": pg.gate.code} if pg.gate else None
    return d


@router.post("", status_code=201)
def create_project(data: ProjectCreate, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    if current_user.role not in ("admin", "editor", "pm", "biz"):
        raise HTTPException(status_code=403, detail="仅管理员、编辑者和商务经理可创建项目")
    data = ProjectCreate(**_fix_empty_strings(data.model_dump()))
    # 类型归一化:空值视为硬件研发项目
    project_type = data.project_type if data.project_type in ("hardware", "software") else "hardware"
    data.project_type = project_type
    # Check unique code
    existing = (db.execute(select(Project).where(Project.code == data.code))).scalar_one_or_none()
    if existing:
        raise HTTPException(status_code=409, detail=f"项目编号 {data.code} 已存在")

    p = Project(**data.model_dump())
    if data.contract_amount is not None:
        p.contract_amount_manual = True
    p.created_at = datetime.utcnow()
    p.updated_at = datetime.utcnow()
    db.add(p)
    db.flush()

    # Auto-create phase gate records for the project's type only (门径评审才能使用)
    phases = (db.execute(select(Phase).where(Phase.project_type == project_type)
                         .order_by(Phase.sort_order))).scalars().all()
    gates = {g.phase_id: g.id for g in db.execute(select(Gate)).scalars().all()}
    gate_names = {g.phase_id: g.name for g in db.execute(select(Gate)).scalars().all()}
    for ph in phases:
        db.add(ProjectPhaseGate(project_id=p.id, phase_id=ph.id, gate_id=gates.get(ph.id), status="not_started"))

    # Auto-create initial milestones (甘特图初始计划)
    if phases:
        ms_base = p.start_date or date.today()
        span_days = (p.planned_end_date - ms_base).days if p.planned_end_date else 60
        span_days = max(span_days, 14)
        if project_type == "software":
            # 软件项目:按阶段必需交付物生成(重要环节粒度,如立项申请书/SRS冻结版)
            from collections import OrderedDict
            stds = (db.execute(
                select(GateDeliverableStandard).join(
                    Phase, Phase.id == GateDeliverableStandard.phase_id
                ).where(
                    Phase.project_type == "software",
                    GateDeliverableStandard.required == True,
                    GateDeliverableStandard.is_active == True,
                ).order_by(Phase.sort_order, GateDeliverableStandard.sort_order)
            )).scalars().all()
            by_phase = OrderedDict()
            for std in stds:
                by_phase.setdefault(std.phase_id, []).append(std)
            phase_days = span_days / max(1, len(by_phase))
            idx_global = 0
            for pidx, (phid, items) in enumerate(by_phase.items()):
                for i, std in enumerate(items):
                    ms_date = ms_base + timedelta(
                        days=int(phase_days * pidx) + int(phase_days * i / max(1, len(items))))
                    db.add(Milestone(
                        project_id=p.id, phase_id=phid, name=std.name,
                        description="按阶段交付物标准自动生成的重要环节",
                        planned_date=ms_date, planned_end_date=ms_date,
                        status="pending", is_key=True, sort_order=idx_global + 1,
                    ))
                    idx_global += 1
        else:
            # 硬件项目:每阶段 1 个门评审里程碑
            interval = max(7, span_days // len(phases))
            for idx, ph in enumerate(phases):
                ms_date = ms_base + timedelta(days=interval * idx)
                ms_name = gate_names.get(ph.id) or f"{ph.name} 评审"
                db.add(Milestone(
                    project_id=p.id, phase_id=ph.id, name=ms_name,
                    description=f"按阶段交付自动生成的初始里程碑（{ph.name}）",
                    planned_date=ms_date,
                    planned_end_date=ms_date + timedelta(days=3),
                    status="pending", is_key=True, sort_order=idx + 1,
                ))

    db.commit()
    db.refresh(p)

    log_audit(db, current_user.username, "create", "project", p.id, p.name, p.id, f"创建项目 {p.code}")
    db.commit()

    # Re-fetch with relationships eagerly loaded
    result3 = db.execute(
        select(Project).where(Project.id == p.id).options(selectinload(Project.client), selectinload(Project.current_phase))
    )
    return project_to_dict(result3.scalar_one())


@router.put("/{project_id}")
def update_project(project_id: int, data: ProjectUpdate, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    result = db.execute(select(Project).where(Project.id == project_id))
    p = result.scalar_one_or_none()
    if not p:
        raise HTTPException(status_code=404, detail="项目不存在")

    upd = data.model_dump(exclude_unset=True)
    _fix_empty_strings(upd)
    upd.pop("version", None)
    upd["version"] = p.version + 1
    upd["updated_at"] = datetime.utcnow()

    # 合同总额手工填写(创建/财务页编辑)后标记,后续交易记录不再覆盖
    if "contract_amount" in upd:
        upd["contract_amount_manual"] = True

    db.execute(update(Project).where(Project.id == project_id).values(**upd))
    db.commit()

    log_audit(db, current_user.username, "update", "project", project_id, p.name, project_id, f"更新项目")
    db.commit()

    # Re-fetch
    result2 = db.execute(select(Project).where(Project.id == project_id).options(
        selectinload(Project.client), selectinload(Project.current_phase), selectinload(Project.members)))
    return project_to_dict(result2.scalars().one())


@router.delete("/{project_id}")
def delete_project(project_id: int, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    result = db.execute(select(Project).where(Project.id == project_id))
    p = result.scalar_one_or_none()
    if not p:
        raise HTTPException(status_code=404, detail="项目不存在")

    db.execute(update(Project).where(Project.id == project_id).values(is_active=False, updated_at=datetime.utcnow()))
    log_audit(db, current_user.username, "delete", "project", project_id, p.name, project_id, f"软删除项目")
    db.commit()
    return {"message": "项目已删除"}


@router.post("/{project_id}/copy", status_code=201)
def copy_project(project_id: int, data: CopyProjectRequest, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    result = db.execute(
        select(Project).where(Project.id == project_id).options(
            selectinload(Project.phase_gates),
            selectinload(Project.milestones),
        )
    )
    src = result.scalars().one_or_none()
    if not src:
        raise HTTPException(status_code=404, detail="源项目不存在")

    existing = (db.execute(select(Project).where(Project.code == data.new_code))).scalar_one_or_none()
    if existing:
        raise HTTPException(status_code=409, detail=f"项目编号 {data.new_code} 已存在")

    new_p = Project(
        name=src.name, code=data.new_code, description=src.description,
        project_type=src.project_type or "hardware",
        client_id=src.client_id, priority=src.priority, difficulty=src.difficulty,
        final_deliverable=src.final_deliverable, acceptance_criteria=src.acceptance_criteria,
        problem_statement=src.problem_statement, aligned_target=src.aligned_target,
        created_at=datetime.utcnow(), updated_at=datetime.utcnow(),
    )
    db.add(new_p)
    db.flush()

    # Copy phase gates (reset to not_started)
    for pg in (src.phase_gates or []):
        db.add(ProjectPhaseGate(
            project_id=new_p.id, phase_id=pg.phase_id, gate_id=pg.gate_id,
            planned_start_date=pg.planned_start_date, planned_end_date=pg.planned_end_date,
            status="not_started",
        ))

    # Copy deliverable templates (reset status)
    dels = (db.execute(select(Deliverable).where(Deliverable.project_id == project_id))).scalars().all()
    for dl in dels:
        db.add(Deliverable(
            project_id=new_p.id, phase_id=dl.phase_id, line_id=dl.line_id,
            name=dl.name, description=dl.description, status="not_submitted",
        ))

    db.commit()
    db.refresh(new_p)

    log_audit(db, current_user.username, "create", "project", new_p.id, new_p.name, new_p.id, f"从 {src.code} 复制")
    db.commit()

    return project_to_dict(new_p)


@router.get("/{project_id}/stats")
def project_stats(project_id: int, db: Session = Depends(get_db), _current_user=Depends(get_current_user)):
    from backend_v2.models import Milestone as MS, RiskIssue as RI, WeeklyReport as WR
    p = (db.execute(select(Project).where(Project.id == project_id))).scalar_one_or_none()
    if not p:
        raise HTTPException(status_code=404, detail="项目不存在")

    ms_total = (db.execute(select(func.count()).select_from(MS).where(MS.project_id == project_id))).scalar()
    ms_done = (db.execute(select(func.count()).select_from(MS).where(MS.project_id == project_id, MS.status == "completed"))).scalar()
    risk_open = (db.execute(select(func.count()).select_from(RI).where(RI.project_id == project_id, RI.status.in_(["open", "in_progress"])))).scalar()
    report_total = (db.execute(select(func.count()).select_from(WR).where(WR.project_id == project_id))).scalar()

    today = date.today()
    cal = today.isocalendar()
    report_this_week = (db.execute(
        select(func.count()).select_from(WR).where(WR.project_id == project_id, WR.year == cal[0], WR.week_number == cal[1])
    )).scalar()

    return {
        "milestone_total": ms_total,
        "milestone_done": ms_done,
        "milestone_rate": round(ms_done / ms_total * 100, 1) if ms_total else 0,
        "open_risks": risk_open,
        "report_total": report_total,
        "report_this_week": report_this_week > 0,
    }
