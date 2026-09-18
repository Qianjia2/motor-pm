"""阶段门评审:交付物标准/检查、任务跟踪与放行、两级签核、历史对比。"""
from datetime import date, datetime
from pathlib import Path
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy import select, func
from sqlalchemy.orm import Session, selectinload
from backend_v2.database import get_db
from backend_v2.auth import get_current_user, require_admin
from backend_v2.config import settings
from backend_v2.permissions import require_permission
from backend_v2.audit import log_audit
from backend_v2.models import (
    Deliverable, Gate, GateDeliverableStandard, GateSignoffRecord,
    Phase, Project, ProjectPhaseGate, ReviewActionItem, TeamMember,
)

router = APIRouter(tags=["gate-reviews"])


def std_to_dict(s):
    d = {c.name: getattr(s, c.name) for c in s.__table__.columns}
    d["phase_name"] = s.phase.name if s.phase else None
    d["phase_code"] = s.phase.code if s.phase else None
    return d


def signoff_to_dict(r):
    d = {c.name: getattr(r, c.name) for c in r.__table__.columns}
    d["signer_name"] = r.signer.name if r.signer else None
    d["signed_at"] = r.signed_at.isoformat() if r.signed_at else None
    d["created_at"] = r.created_at.isoformat() if r.created_at else None
    d["gate_id"] = r.phase_gate.gate_id if r.phase_gate else None
    d["phase_id"] = r.phase_gate.phase_id if r.phase_gate else None
    d["initiated_by"] = r.phase_gate.signoff_initiated_by if r.phase_gate else None
    return d


def pg_brief(pg):
    return {
        "id": pg.id, "project_id": pg.project_id,
        "project_name": pg.project.name if pg.project else None,
        "phase_id": pg.phase_id,
        "phase_name": pg.phase.name if pg.phase else None,
        "phase_code": pg.phase.code if pg.phase else None,
        "gate_id": pg.gate_id,
        "gate_name": pg.gate.name if pg.gate else None,
        "gate_code": pg.gate.code if pg.gate else None,
        "gate_status": pg.gate_status,
        "status": pg.status,
        "gate_review_date": pg.gate_review_date.isoformat() if pg.gate_review_date else None,
        "gate_review_notes": pg.gate_review_notes,
        "planned_end_date": pg.planned_end_date.isoformat() if pg.planned_end_date else None,
    }


# ── 交付物标准 CRUD(管理员) ──

@router.get("/api/gate-deliverable-standards")
def list_standards(
    phase_id: int = Query(None),
    project_type: str = Query(None),
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
):
    q = select(GateDeliverableStandard).where(GateDeliverableStandard.is_active == True)
    if phase_id:
        q = q.where(GateDeliverableStandard.phase_id == phase_id)
    if project_type:
        q = q.join(Phase, Phase.id == GateDeliverableStandard.phase_id).where(Phase.project_type == project_type)
    q = q.order_by(GateDeliverableStandard.phase_id, GateDeliverableStandard.sort_order, GateDeliverableStandard.id)
    return [std_to_dict(s) for s in db.execute(q).scalars().all()]


@router.post("/api/gate-deliverable-standards", status_code=201)
def create_standard(
    data: dict, db: Session = Depends(get_db), _user=Depends(require_admin),
):
    if not data.get("phase_id") or not data.get("name", "").strip():
        raise HTTPException(status_code=400, detail="阶段和名称必填")
    s = GateDeliverableStandard(
        phase_id=int(data["phase_id"]), name=data["name"].strip(),
        category=(data.get("category") or "document")[:64],
        required=bool(data.get("required", True)),
        description=data.get("description") or None,
        acceptance_criteria=data.get("acceptance_criteria") or None,
        responsible_role=(data.get("responsible_role") or None)[:64] if data.get("responsible_role") else None,
        sort_order=int(data.get("sort_order") or 0),
    )
    db.add(s); db.commit(); db.refresh(s)
    log_audit(db, _user.username, "create", "gate_deliverable_standard", s.id, s.name, None, "新增交付物标准")
    return std_to_dict(s)


@router.put("/api/gate-deliverable-standards/{std_id}")
def update_standard(
    std_id: int, data: dict, db: Session = Depends(get_db), _user=Depends(require_admin),
):
    s = db.get(GateDeliverableStandard, std_id)
    if not s:
        raise HTTPException(status_code=404, detail="标准不存在")
    allowed = {"phase_id", "name", "category", "required", "description", "acceptance_criteria", "responsible_role", "sort_order", "is_active"}
    for k, v in data.items():
        if k in allowed and v is not None:
            if k in ("phase_id", "sort_order"):
                setattr(s, k, int(v))
            elif k == "required":
                setattr(s, k, bool(v))
            else:
                setattr(s, k, v)
    db.commit(); db.refresh(s)
    log_audit(db, _user.username, "update", "gate_deliverable_standard", s.id, s.name, None, "更新交付物标准")
    return std_to_dict(s)


@router.delete("/api/gate-deliverable-standards/{std_id}")
def delete_standard(
    std_id: int, db: Session = Depends(get_db), _user=Depends(require_admin),
):
    s = db.get(GateDeliverableStandard, std_id)
    if not s:
        raise HTTPException(status_code=404, detail="标准不存在")
    # 软删而不是硬删:每次启动 init_db() 都会跑播种(_seed_*_standards),
    # 它按 (phase_id, name)「查不到就插」,硬删掉的行会被当成"从没存在过"原样插回来。
    # 留着行置 is_active=0,查重才命中并跳过。已停用的再删直接返回,不重复写审计。
    if s.is_active:
        s.is_active = False
        db.commit()
        log_audit(db, _user.username, "delete", "gate_deliverable_standard", std_id, s.name, None, "停用交付物标准")
    return {"ok": True}


@router.get("/api/gate-deliverable-standards/{std_id}/template")
def standard_template(
    std_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user),
):
    """下载交付物标准模板(Word),按分类生成填写框架。"""
    s = db.get(GateDeliverableStandard, std_id)
    if not s:
        raise HTTPException(status_code=404, detail="标准不存在")
    if s.template_path:  # 已上传自定义模板则直接返回
        from fastapi.responses import FileResponse
        tpl_file = Path(settings.UPLOAD_DIR) / s.template_path
        if tpl_file.exists():
            return FileResponse(str(tpl_file), filename=tpl_file.name)
    from io import BytesIO
    from urllib.parse import quote
    from fastapi.responses import StreamingResponse
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    SECTIONS = {
        "document": ["目的与范围", "正文内容", "参考依据", "附录(图纸/数据/协议)"],
        "report": ["背景与目的", "方法与过程", "结果与结论", "问题与建议"],
        "approval": ["审批事项", "审批依据", "审批意见", "审批人签字/日期"],
        "other": ["内容", "结论", "备注"],
    }
    cat = s.category or "other"
    cat_name = {"document": "文档", "report": "报告", "approval": "审批", "other": "其他"}.get(cat, "其他")

    doc = Document()
    doc.styles["Normal"].font.name = "SimSun"
    doc.styles["Normal"].font.size = Pt(11)
    for sec in doc.sections:
        sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)
        sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)

    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(s.name); r.bold = True; r.font.size = Pt(16)
    doc.add_paragraph()

    t = doc.add_table(rows=0, cols=4); t.style = "Table Grid"
    for a, b, c, d in [
        ("阶段", s.phase.name if s.phase else "—", "分类", cat_name),
        ("是否必需", "必需" if s.required else "可选", "责任角色", s.responsible_role or "—"),
    ]:
        cells = t.add_row().cells
        cells[0].text = a; cells[1].text = b; cells[2].text = c; cells[3].text = d
    doc.add_paragraph()

    def label(text):
        p = doc.add_paragraph()
        p.add_run(text).bold = True
        return p

    if s.description:
        label("交付物描述")
        doc.add_paragraph(s.description)
    if s.acceptance_criteria:
        label("通过标准")
        doc.add_paragraph(s.acceptance_criteria)
    doc.add_paragraph()

    for sec in SECTIONS.get(cat, SECTIONS["other"]):
        label(sec)
        for _ in range(3):
            doc.add_paragraph("")
    if cat == "approval":
        doc.add_table(rows=3, cols=4).style = "Table Grid"

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    fname = f"模板-{s.name}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"},
    )


@router.post("/api/gate-deliverable-standards/{std_id}/template")
async def upload_standard_template(
    std_id: int, file: UploadFile = File(...),
    db: Session = Depends(get_db), _user=Depends(require_admin),
):
    """上传交付物标准模板文件(替换旧文件)。"""
    s = db.get(GateDeliverableStandard, std_id)
    if not s:
        raise HTTPException(status_code=404, detail="标准不存在")
    ext = file.filename.rsplit(".", 1)[-1].lower() if file.filename else ""
    if ext not in {"doc", "docx", "xls", "xlsx", "ppt", "pptx", "pdf", "zip"}:
        raise HTTPException(status_code=400, detail=f"不支持的模板类型: .{ext}")
    content = await file.read()
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="文件不超过 20MB")
    from uuid import uuid4
    tpl_dir = Path(settings.UPLOAD_DIR) / "gate_templates"
    tpl_dir.mkdir(parents=True, exist_ok=True)
    dest = tpl_dir / f"std{std_id}_{uuid4().hex[:8]}_{file.filename}"
    with open(dest, "wb") as f:
        f.write(content)
    if s.template_path:  # 删除旧模板
        try:
            (Path(settings.UPLOAD_DIR) / s.template_path).unlink(missing_ok=True)
        except Exception:
            pass
    s.template_path = str(dest.relative_to(Path(settings.UPLOAD_DIR)))
    db.commit()
    log_audit(db, _user.username, "update", "gate_deliverable_standard", s.id, s.name, None, "上传交付物模板")
    return {"message": "模板已上传", "template_path": s.template_path}


@router.delete("/api/gate-deliverable-standards/{std_id}/template")
def delete_standard_template(
    std_id: int, db: Session = Depends(get_db), _user=Depends(require_admin),
):
    """移除已上传的模板,恢复为自动生成骨架。"""
    s = db.get(GateDeliverableStandard, std_id)
    if not s:
        raise HTTPException(status_code=404, detail="标准不存在")
    if s.template_path:
        try:
            (Path(settings.UPLOAD_DIR) / s.template_path).unlink(missing_ok=True)
        except Exception:
            pass
        s.template_path = None
        db.commit()
        log_audit(db, _user.username, "update", "gate_deliverable_standard", s.id, s.name, None, "删除交付物模板")
    return {"ok": True}


# ── 对照检查:标准 x 各项目实际交付物 ──

@router.get("/api/gate-reviews/checklist")
def deliverable_checklist(
    phase_id: int = Query(..., ge=1),
    project_id: int = Query(None),
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
):
    standards = db.execute(
        select(GateDeliverableStandard).where(GateDeliverableStandard.phase_id == phase_id, GateDeliverableStandard.is_active == True)
        .order_by(GateDeliverableStandard.sort_order, GateDeliverableStandard.id)
    ).scalars().all()

    # 有该阶段门的正式项目(可按项目过滤);已删除项目(is_active=False)不参与对照
    q_pg = select(ProjectPhaseGate).join(Project).options(
        selectinload(ProjectPhaseGate.project), selectinload(ProjectPhaseGate.phase), selectinload(ProjectPhaseGate.gate)
    ).where(ProjectPhaseGate.phase_id == phase_id, Project.is_active == True)
    if project_id:
        q_pg = q_pg.where(ProjectPhaseGate.project_id == project_id)
    pgs = db.execute(q_pg).scalars().all()

    delivs = db.execute(
        select(Deliverable).options(selectinload(Deliverable.project)).where(Deliverable.phase_id == phase_id)
    ).scalars().all()

    items = []
    for st in standards:
        proj_items = []
        for pg in pgs:
            pid = pg.project_id
            pname = pg.project.name if pg.project else f"项目{pid}"
            # 一个标准可关联多个交付物:按状态优先级排序(approved 最优),全部返回
            match = sorted(
                (d for d in delivs if d.project_id == pid and d.name and st.name and (d.name in st.name or st.name in d.name)),
                key=lambda d: {"approved": 0, "submitted": 1, "reviewed": 2, "rejected": 3, "not_submitted": 4}.get(d.status, 5),
            )
            if match:
                m = match[0]
                proj_items.append({
                    "project_id": pid, "project_name": pname,
                    "deliverable_id": m.id, "deliverable_name": m.name,
                    "status": m.status,
                    "submitted_date": m.submitted_date.isoformat() if m.submitted_date else None,
                    "owner_name": m.owner.name if m.owner else None,
                    "deliverables": [{
                        "deliverable_id": d.id, "deliverable_name": d.name,
                        "status": d.status,
                        "submitted_date": d.submitted_date.isoformat() if d.submitted_date else None,
                        "owner_name": d.owner.name if d.owner else None,
                        "file_path": d.file_path,
                    } for d in match],
                })
            else:
                proj_items.append({
                    "project_id": pid, "project_name": pname,
                    "deliverable_id": None, "deliverable_name": None,
                    "status": "missing", "submitted_date": None, "owner_name": None,
                    "deliverables": [],
                })
        items.append({"id": st.id, "name": st.name, "category": st.category,
                      "required": st.required, "description": st.description,
                      "acceptance_criteria": st.acceptance_criteria, "responsible_role": st.responsible_role,
                      "sort_order": st.sort_order,
                      "projects": proj_items})

    return {
        "phase_id": phase_id,
        "phase_name": pgs[0].phase.name if pgs else (db.get(Phase, phase_id).name if db.get(Phase, phase_id) else None),
        "standards": items,
    }


# ── 任务跟踪(跨项目) ──

@router.get("/api/gate-reviews/action-items")
def cross_project_action_items(
    project_id: int = Query(None),
    phase_id: int = Query(None),
    status: str = Query(None),
    assignee_id: int = Query(None),
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
):
    q = select(ReviewActionItem).options(
        selectinload(ReviewActionItem.phase_gate).selectinload(ProjectPhaseGate.project),
        selectinload(ReviewActionItem.phase_gate).selectinload(ProjectPhaseGate.phase),
        selectinload(ReviewActionItem.phase_gate).selectinload(ProjectPhaseGate.gate),
        selectinload(ReviewActionItem.assignee),
    ).join(ProjectPhaseGate)
    if project_id:
        q = q.where(ProjectPhaseGate.project_id == project_id)
    if phase_id:
        q = q.where(ProjectPhaseGate.phase_id == phase_id)
    if status:
        q = q.where(ReviewActionItem.status == status)
    if assignee_id:
        q = q.where(ReviewActionItem.assignee_id == assignee_id)
    q = q.order_by(ReviewActionItem.status == "closed", ReviewActionItem.due_date, ReviewActionItem.id)

    rows = []
    for ai in db.execute(q).scalars().all():
        pg = ai.phase_gate
        rows.append({
            "id": ai.id, "title": ai.title, "description": ai.description,
            "status": ai.status, "due_date": ai.due_date.isoformat() if ai.due_date else None,
            "resolution": ai.resolution, "created_at": ai.created_at.isoformat() if ai.created_at else None,
            "closed_at": ai.closed_at.isoformat() if ai.closed_at else None,
            "assignee_id": ai.assignee_id, "assignee_name": ai.assignee.name if ai.assignee else None,
            "project_phase_gate_id": pg.id if pg else None,
            "project_id": pg.project_id if pg else None,
            "project_name": pg.project.name if pg and pg.project else None,
            "phase_id": pg.phase_id if pg else None,
            "phase_name": pg.phase.name if pg and pg.phase else None,
            "gate_id": pg.gate_id if pg else None,
            "gate_name": pg.gate.name if pg and pg.gate else None,
            "gate_status": pg.gate_status if pg else None,
        })
    return rows


# ── 发起放行 + 两级签核 ──

@router.post("/api/gate-reviews/{pg_id}/initiate-signoff", status_code=201)
def initiate_signoff(
    pg_id: int, data: dict,
    db: Session = Depends(get_db), _user=Depends(require_permission("projects", "edit")),
):
    pg = db.execute(
        select(ProjectPhaseGate).options(selectinload(ProjectPhaseGate.project), selectinload(ProjectPhaseGate.phase)).where(ProjectPhaseGate.id == pg_id)
    ).scalar_one_or_none()
    if not pg:
        raise HTTPException(status_code=404, detail="阶段门不存在")
    if pg.gate_status == "passed":
        raise HTTPException(status_code=400, detail="该门已放行,无需重复发起")
    existing = db.execute(
        select(GateSignoffRecord).where(GateSignoffRecord.project_phase_gate_id == pg_id,
                                        GateSignoffRecord.status == "pending")
    ).scalars().all()
    if existing:
        raise HTTPException(status_code=400, detail="该门已有进行中的签核流程")
    missing = _missing_required_deliverables(db, pg)
    if missing:
        raise HTTPException(status_code=400, detail="该门尚有必需交付物未批准,暂不能发起签核: " + "、".join(missing))
    l1 = data.get("level1_signer_id"); l2 = data.get("level2_signer_id")
    if not l1 or not l2:
        raise HTTPException(status_code=400, detail="请指定两级签核人")
    if not db.get(TeamMember, int(l1)) or not db.get(TeamMember, int(l2)):
        raise HTTPException(status_code=400, detail="签核人不存在")
    pg.signoff_initiated_by = _user.username
    db.add_all([
        GateSignoffRecord(project_phase_gate_id=pg_id, level=1, signer_id=int(l1)),
        GateSignoffRecord(project_phase_gate_id=pg_id, level=2, signer_id=int(l2)),
    ])
    db.commit()
    log_audit(db, _user.username, "create", "phase_gate", pg_id,
              f"{pg.phase.name if pg.phase else ''}门评审放行", pg.project_id, "发起阶段门放行签核")
    recs = db.execute(
        select(GateSignoffRecord).where(GateSignoffRecord.project_phase_gate_id == pg_id).order_by(GateSignoffRecord.level)
    ).scalars().all()
    return [signoff_to_dict(r) for r in recs]


def _signoff_permission_check(rec, current_user):
    """签核人本人或管理员可操作。"""
    if current_user.role == "admin":
        return
    if current_user.member_id and current_user.member_id == rec.signer_id:
        return
    raise HTTPException(status_code=403, detail="仅签核人本人或管理员可操作")


@router.post("/api/gate-reviews/{pg_id}/withdraw")
def withdraw_gate_signoff(
    pg_id: int,
    db: Session = Depends(get_db), current_user=Depends(get_current_user),
):
    """发起人本人撤回阶段门签核流程:删除该门全部签核记录并重置放行状态,可重新发起。"""
    pg = db.execute(
        select(ProjectPhaseGate).options(selectinload(ProjectPhaseGate.project), selectinload(ProjectPhaseGate.phase)).where(ProjectPhaseGate.id == pg_id)
    ).scalar_one_or_none()
    if not pg:
        raise HTTPException(status_code=404, detail="阶段门不存在")
    recs = db.execute(
        select(GateSignoffRecord).where(GateSignoffRecord.project_phase_gate_id == pg_id)
    ).scalars().all()
    if not any(r.status == "pending" for r in recs):
        raise HTTPException(status_code=400, detail="该门无进行中的签核流程")
    if current_user.role != "admin" and pg.signoff_initiated_by != current_user.username:
        raise HTTPException(status_code=403, detail="仅签核发起人本人或管理员可操作")
    for r in recs:
        db.delete(r)
    pg.signoff_initiated_by = None
    if pg.gate_status in ("passed", "failed"):
        pg.gate_status = None
        pg.gate_review_date = None
    db.commit()
    log_audit(db, current_user.username, "delete", "phase_gate", pg_id,
              f"{pg.phase.name if pg.phase else ''}门签核流程撤回", pg.project_id,
              f"发起人撤回签核流程,删除记录 {len(recs)} 条")
    return {"withdrawn": len(recs)}


def _missing_required_deliverables(db, pg):
    """该阶段门未达标的必需标准交付物名称列表(名称双向匹配,需 approved)。"""
    stds = db.execute(
        select(GateDeliverableStandard).where(
            GateDeliverableStandard.phase_id == pg.phase_id,
            GateDeliverableStandard.required == True,
            GateDeliverableStandard.is_active == True,
        )
    ).scalars().all()
    if not stds:
        return []
    dels = db.execute(select(Deliverable).where(
        Deliverable.project_id == pg.project_id,
        Deliverable.phase_id == pg.phase_id,
    )).scalars().all()
    missing = []
    for st in stds:
        if not st.name:
            continue
        ok = False
        for dl in dels:
            if dl.name and dl.status == "approved" and (st.name in dl.name or dl.name in st.name):
                ok = True
                break
        if not ok:
            missing.append(st.name)
    return missing


@router.get("/api/gate-reviews/my-signoffs")
def my_signoffs(db: Session = Depends(get_db), _user=Depends(get_current_user)):
    if not _user.member_id:
        return []
    q = select(GateSignoffRecord).options(
        selectinload(GateSignoffRecord.phase_gate).selectinload(ProjectPhaseGate.project),
        selectinload(GateSignoffRecord.phase_gate).selectinload(ProjectPhaseGate.phase),
        selectinload(GateSignoffRecord.phase_gate).selectinload(ProjectPhaseGate.gate),
    ).where(GateSignoffRecord.signer_id == _user.member_id, GateSignoffRecord.status == "pending")
    rows = []
    for r in db.execute(q.order_by(GateSignoffRecord.created_at)).scalars().all():
        d = signoff_to_dict(r)
        d["project_id"] = r.phase_gate.project_id if r.phase_gate else None
        d["project_name"] = r.phase_gate.project.name if r.phase_gate and r.phase_gate.project else None
        d["phase_name"] = r.phase_gate.phase.name if r.phase_gate and r.phase_gate.phase else None
        d["gate_name"] = r.phase_gate.gate.name if r.phase_gate and r.phase_gate.gate else None
        rows.append(d)
    return rows


@router.get("/api/gate-reviews/signoffs")
def all_signoffs(
    project_id: int = Query(None),
    status: str = Query(None),
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
):
    q = select(GateSignoffRecord).options(
        selectinload(GateSignoffRecord.phase_gate).selectinload(ProjectPhaseGate.project),
        selectinload(GateSignoffRecord.phase_gate).selectinload(ProjectPhaseGate.phase),
        selectinload(GateSignoffRecord.phase_gate).selectinload(ProjectPhaseGate.gate),
    )
    if project_id:
        q = q.join(ProjectPhaseGate).where(ProjectPhaseGate.project_id == project_id)
    if status:
        q = q.where(GateSignoffRecord.status == status)
    rows = []
    for r in db.execute(q.order_by(GateSignoffRecord.created_at.desc())).scalars().all():
        d = signoff_to_dict(r)
        d["project_id"] = r.phase_gate.project_id if r.phase_gate else None
        d["project_name"] = r.phase_gate.project.name if r.phase_gate and r.phase_gate.project else None
        d["phase_name"] = r.phase_gate.phase.name if r.phase_gate and r.phase_gate.phase else None
        d["gate_name"] = r.phase_gate.gate.name if r.phase_gate and r.phase_gate.gate else None
        rows.append(d)
    return rows


@router.post("/api/gate-reviews/signoffs/{signoff_id}/approve")
def approve_signoff(
    signoff_id: int, data: dict,
    db: Session = Depends(get_db), current_user=Depends(require_permission("projects", "edit")),
):
    rec = db.get(GateSignoffRecord, signoff_id)
    if not rec:
        raise HTTPException(status_code=404, detail="签核记录不存在")
    _signoff_permission_check(rec, current_user)
    if rec.status != "pending":
        raise HTTPException(status_code=400, detail="该记录已处理")
    pg = db.get(ProjectPhaseGate, rec.project_phase_gate_id)
    if not pg:
        raise HTTPException(status_code=404, detail="阶段门不存在")
    if rec.level == 2:
        l1 = db.execute(
            select(GateSignoffRecord).where(GateSignoffRecord.project_phase_gate_id == pg.id,
                                            GateSignoffRecord.level == 1)
        ).scalar_one_or_none()
        if not l1 or l1.status != "approved":
            raise HTTPException(status_code=400, detail="需先完成一级签核")
    # 两级通过均校验:该阶段必需交付物必须已批准
    missing = _missing_required_deliverables(db, pg)
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"该门尚有必需交付物未批准,暂不能{'放行' if rec.level == 2 else '通过'}: " + "、".join(missing),
        )
    if rec.level == 2:
        pg.gate_status = "passed"
        pg.gate_review_date = date.today()
    rec.status = "approved"
    rec.comment = (data.get("comment") or "").strip() or None
    rec.signed_by_username = current_user.username
    rec.signed_at = datetime.utcnow()
    db.commit()
    log_audit(db, current_user.username, "update", "phase_gate", pg.id,
              f"{pg.phase.name if pg.phase else ''}门签核通过(L{rec.level})", pg.project_id,
              "签核通过" + (f":{rec.comment}" if rec.comment else ""))
    return signoff_to_dict(rec)


@router.post("/api/gate-reviews/signoffs/{signoff_id}/reject")
def reject_signoff(
    signoff_id: int, data: dict,
    db: Session = Depends(get_db), current_user=Depends(require_permission("projects", "edit")),
):
    rec = db.get(GateSignoffRecord, signoff_id)
    if not rec:
        raise HTTPException(status_code=404, detail="签核记录不存在")
    _signoff_permission_check(rec, current_user)
    if rec.status != "pending":
        raise HTTPException(status_code=400, detail="该记录已处理")
    pg = db.get(ProjectPhaseGate, rec.project_phase_gate_id)
    if not pg:
        raise HTTPException(status_code=404, detail="阶段门不存在")
    rec.status = "rejected"
    rec.comment = (data.get("comment") or "").strip() or "驳回"
    rec.signed_by_username = current_user.username
    rec.signed_at = datetime.utcnow()
    pg.gate_status = "failed"
    db.commit()
    log_audit(db, current_user.username, "update", "phase_gate", pg.id,
              f"{pg.phase.name if pg.phase else ''}门签核驳回(L{rec.level})", pg.project_id,
              f"签核驳回:{rec.comment}")
    return signoff_to_dict(rec)


@router.delete("/api/gate-reviews/signoffs/{signoff_id}")
def delete_signoff(
    signoff_id: int,
    db: Session = Depends(get_db), current_user=Depends(require_admin),
):
    """仅管理员可删除:撤销该门发起的签核流程(删除该门全部签核记录),已放行/未通过的门重置为未放行。"""
    rec = db.get(GateSignoffRecord, signoff_id)
    if not rec:
        raise HTTPException(status_code=404, detail="签核记录不存在")
    pg = db.get(ProjectPhaseGate, rec.project_phase_gate_id)
    recs = db.execute(
        select(GateSignoffRecord).where(GateSignoffRecord.project_phase_gate_id == rec.project_phase_gate_id)
    ).scalars().all()
    for r in recs:
        db.delete(r)
    if pg:
        pg.signoff_initiated_by = None
        if pg.gate_status in ("passed", "failed"):
            pg.gate_status = None
            pg.gate_review_date = None
    db.commit()
    log_audit(db, current_user.username, "delete", "phase_gate", rec.project_phase_gate_id,
              f"{pg.phase.name if pg and pg.phase else ''}门签核流程删除" if pg else "签核记录删除",
              pg.project_id if pg else None, f"管理员删除签核记录 {len(recs)} 条")
    return {"deleted": len(recs)}


# ── 历史对比 ──

@router.get("/api/gate-reviews/history")
def gate_history(
    view: str = Query("by_gate", pattern="^(by_gate|by_project)$"),
    gate_id: int = Query(None),
    project_id: int = Query(None),
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
):
    pgs = db.execute(
        select(ProjectPhaseGate).options(
            selectinload(ProjectPhaseGate.project), selectinload(ProjectPhaseGate.phase),
            selectinload(ProjectPhaseGate.gate), selectinload(ProjectPhaseGate.signoff_records).selectinload(GateSignoffRecord.signer),
        )
    ).scalars().all()
    # 仅统计启用中的项目,避免停用/测试残留项目出现在历史对比中
    pgs = [p for p in pgs if p.project and p.project.is_active]

    if view == "by_gate":
        gates = db.execute(select(Gate).options(selectinload(Gate.phase)).order_by(Gate.phase_id, Gate.sort_order)).scalars().all()
        out = []
        for g in gates:
            if gate_id and g.id != gate_id:
                continue
            g_pgs = [p for p in pgs if p.gate_id == g.id and (not project_id or p.project_id == project_id)]
            passed = sum(1 for p in g_pgs if p.gate_status == "passed")
            out.append({
                "gate_id": g.id, "gate_name": g.name, "gate_code": g.code,
                "phase_name": g.phase.name if g.phase else None,
                "total": len(g_pgs), "passed": passed,
                "pass_rate": round(passed / len(g_pgs), 2) if g_pgs else None,
                "projects": [{
                    "project_id": p.project_id, "project_name": p.project.name if p.project else None,
                    "gate_status": p.gate_status, "status": p.status,
                    "review_date": p.gate_review_date.isoformat() if p.gate_review_date else None,
                    "notes": p.gate_review_notes,
                    "signoffs": [{
                        "level": s.level, "status": s.status,
                        "signer_name": s.signer.name if s.signer else None,
                        "signed_at": s.signed_at.isoformat() if s.signed_at else None,
                        "comment": s.comment,
                    } for s in (p.signoff_records or [])],
                } for p in g_pgs],
            })
        return {"view": "by_gate", "gates": out}

    projects = db.execute(select(Project).where(Project.is_active == 1).order_by(Project.id)).scalars().all()
    out = []
    for pr in projects:
        if project_id and pr.id != project_id:
            continue
        pr_pgs = [p for p in pgs if p.project_id == pr.id and (not gate_id or p.gate_id == gate_id)]
        out.append({
            "project_id": pr.id, "project_name": pr.name, "code": pr.code,
            "total": len(pr_pgs),
            "passed": sum(1 for p in pr_pgs if p.gate_status == "passed"),
            "gates": [{
                "gate_id": p.gate_id, "gate_name": p.gate.name if p.gate else None,
                "gate_code": p.gate.code if p.gate else None,
                "phase_name": p.phase.name if p.phase else None,
                "gate_status": p.gate_status, "status": p.status,
                "review_date": p.gate_review_date.isoformat() if p.gate_review_date else None,
                "planned_end_date": p.planned_end_date.isoformat() if p.planned_end_date else None,
                "signoffs": [{
                    "level": s.level, "status": s.status,
                    "signer_name": s.signer.name if s.signer else None,
                    "signed_at": s.signed_at.isoformat() if s.signed_at else None,
                    "comment": s.comment,
                } for s in (p.signoff_records or [])],
            } for p in pr_pgs],
        })
    return {"view": "by_project", "projects": out}
