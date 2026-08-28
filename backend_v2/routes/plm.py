"""PLM: BOM + ECN management."""
import json, os
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form
from fastapi.responses import FileResponse
from sqlalchemy import select
from sqlalchemy.orm import Session
from backend_v2.database import get_db
from backend_v2.auth import get_current_user
from backend_v2.models import BomItem, Ecn, Project
from backend_v2.storage import safe_load, safe_save
import openpyxl
from io import BytesIO

router = APIRouter(tags=["plm"])

CENTRAL_BOM_FILE = os.path.join(os.path.dirname(__file__), "..", "..", "data", "bom.json")


def _sync_to_central_bom(item: dict, project_id: int, action: str):
    """Sync project BOM item to central BOM JSON. action: create/update/delete."""
    try:
        items = []
        if os.path.exists(CENTRAL_BOM_FILE):
            with open(CENTRAL_BOM_FILE, "r", encoding="utf-8") as f:
                items = json.load(f)

        source_key = f"proj-{project_id}-{item.get('part_number','')}-{item.get('name','')}"

        if action == "delete":
            items = [x for x in items if x.get("_source") != source_key]
        else:
            # create or update — upsert by source_key
            existing = next((x for x in items if x.get("_source") == source_key), None)
            spec_parts = [item.get("specification", ""), item.get("material", "")]
            spec = " / ".join(p for p in spec_parts if p) or ""
            entry = {
                "id": existing["id"] if existing else (max([x.get("id", 0) for x in items], default=0) + 1),
                "category": "电机类物料",
                "code": item.get("part_number", ""),
                "name": item.get("name", ""),
                "spec": spec,
                "unit": item.get("unit", "pcs"),
                "qty": item.get("quantity", 1),
                "supplier": item.get("supplier", ""),
                "supplier_code": "",
                "price": "",
                "notes": item.get("notes", ""),
                "_source": source_key,
                "_project_id": project_id,
                "_bom_id": item.get("id"),
                "created_at": item.get("created_at") or datetime.utcnow().isoformat(),
            }
            if existing:
                items = [entry if x.get("_source") == source_key else x for x in items]
            else:
                items.append(entry)

        with open(CENTRAL_BOM_FILE, "w", encoding="utf-8") as f:
            json.dump(items, f, ensure_ascii=False, indent=2)
    except Exception:
        pass  # silent sync failure, don't block BOM editing


# ═══════════════════ BOM ═══════════════════

def bom_to_dict(item):
    d = {c.name: getattr(item, c.name) for c in item.__table__.columns}
    return d


def build_bom_tree(items, parent_id=None):
    children = [i for i in items if i.get("parent_id") == parent_id]
    children.sort(key=lambda x: (x.get("sort_order", 0), x.get("id", 0)))
    for c in children:
        c["children"] = build_bom_tree(items, c["id"])
    return children


@router.get("/api/projects/{project_id}/bom")
def list_bom(project_id: int, build_id: int = Query(None), db: Session = Depends(get_db), _user=Depends(get_current_user)):
    q = select(BomItem).where(BomItem.project_id == project_id)
    if build_id:
        q = q.where(BomItem.build_id == build_id)
    else:
        q = q.where(BomItem.build_id.is_(None))
    result = db.execute(q.order_by(BomItem.sort_order, BomItem.id))
    items = [bom_to_dict(i) for i in result.scalars().all()]
    return build_bom_tree(items)


@router.post("/api/projects/{project_id}/bom", status_code=201)
def create_bom_item(project_id: int, data: dict, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    allowed = {"parent_id", "level", "part_number", "name", "specification", "material",
               "quantity", "unit", "supplier", "drawing_file", "model_file", "notes", "sort_order", "build_id"}
    item = BomItem(project_id=project_id, **{k: v for k, v in data.items() if k in allowed})
    db.add(item); db.commit(); db.refresh(item)
    result = bom_to_dict(item)
    _sync_to_central_bom(result, project_id, "create")
    return result


@router.put("/api/bom/{item_id}")
def update_bom_item(item_id: int, data: dict, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    from sqlalchemy import update as up
    r = (db.execute(select(BomItem).where(BomItem.id == item_id))).scalar_one_or_none()
    if not r: raise HTTPException(status_code=404, detail="不存在")
    allowed = {"parent_id", "level", "part_number", "name", "specification", "material",
               "quantity", "unit", "supplier", "drawing_file", "model_file", "notes", "sort_order"}
    upd = {k: v for k, v in data.items() if k in allowed}
    if upd: db.execute(up(BomItem).where(BomItem.id == item_id).values(**upd)); db.commit()
    r2 = (db.execute(select(BomItem).where(BomItem.id == item_id))).scalar_one_or_none()
    if r2: _sync_to_central_bom(bom_to_dict(r2), r2.project_id, "update")
    return {"message": "已更新"}


@router.delete("/api/bom/{item_id}")
def delete_bom_item(item_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    from sqlalchemy import delete as dl, update as up
    r = (db.execute(select(BomItem).where(BomItem.id == item_id))).scalar_one_or_none()
    if not r: raise HTTPException(status_code=404, detail="不存在")
    item_dict = bom_to_dict(r)
    db.execute(up(BomItem).where(BomItem.parent_id == item_id).values(parent_id=r.parent_id))
    db.execute(dl(BomItem).where(BomItem.id == item_id)); db.commit()
    _sync_to_central_bom(item_dict, r.project_id, "delete")
    return {"message": "已删除"}


# ═══════════════════ ECN ═══════════════════

@router.get("/api/projects/{project_id}/ecns")
def list_ecns(project_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    result = db.execute(select(Ecn).where(Ecn.project_id == project_id).order_by(Ecn.created_at.desc()))
    return [{c.name: getattr(e, c.name) for c in e.__table__.columns} for e in result.scalars().all()]


@router.post("/api/projects/{project_id}/ecns", status_code=201)
def create_ecn(project_id: int, data: dict, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    from datetime import date as date_type
    def _pd(d, f):  # parse date
        if f in d and d[f] and isinstance(d[f], str):
            try: d[f] = date_type.fromisoformat(d[f])
            except: pass
    _pd(data, "submitted_date")
    # Auto-generate ECN number
    count = (db.execute(select(Ecn).where(Ecn.project_id == project_id))).scalars().all()
    ecn_num = f"ECN-{project_id:04d}-{len(count)+1:03d}"
    allowed = {"title", "reason", "change_description", "affected_items", "status", "submitted_by", "submitted_date"}
    e = Ecn(project_id=project_id, ecn_number=ecn_num, created_at=datetime.utcnow(),
            **{k: v for k, v in data.items() if k in allowed})
    db.add(e); db.commit(); db.refresh(e)
    return {c.name: getattr(e, c.name) for c in e.__table__.columns}


@router.put("/api/ecns/{ecn_id}")
def update_ecn(ecn_id: int, data: dict, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    from sqlalchemy import update as up
    from datetime import date as date_type
    def _pd(d, f):
        if f in d and d[f] and isinstance(d[f], str):
            try: d[f] = date_type.fromisoformat(d[f])
            except: pass
    _pd(data, "submitted_date"); _pd(data, "approved_date")
    r = (db.execute(select(Ecn).where(Ecn.id == ecn_id))).scalar_one_or_none()
    if not r: raise HTTPException(status_code=404, detail="不存在")
    allowed = {"title", "reason", "change_description", "affected_items", "status", "submitted_by", "approved_by", "submitted_date", "approved_date"}
    upd = {k: v for k, v in data.items() if k in allowed}
    if upd: db.execute(up(Ecn).where(Ecn.id == ecn_id).values(**upd)); db.commit()
    return {"message": "已更新"}


@router.delete("/api/ecns/{ecn_id}")
def delete_ecn(ecn_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    from sqlalchemy import delete as dl
    db.execute(dl(Ecn).where(Ecn.id == ecn_id)); db.commit()
    return {"message": "已删除"}


# ═══════════════════ BOM Excel Upload ═══════════════════

@router.post("/api/projects/{project_id}/bom/upload")
def upload_bom(project_id: int, file: UploadFile = File(...),
                     build_id: int = Query(None),
                     db: Session = Depends(get_db), _user=Depends(get_current_user)):
    """Upload BOM from any format — auto-detect, no fixed template required."""
    import re, io, csv
    content = file.file.read()
    all_rows = []

    # Try Excel parsers first
    excel_ok = False
    for parser_name, parser_fn in [
        ("openpyxl", lambda c: openpyxl.load_workbook(io.BytesIO(c)).active),
        ("xlrd", lambda c: __import__("xlrd").open_workbook(file_contents=c).sheet_by_index(0)),
    ]:
        try:
            sheet = parser_fn(content)
            if parser_name == "openpyxl":
                all_rows = list(sheet.iter_rows(values_only=True))
            else:
                all_rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
            excel_ok = True
            break
        except:
            continue

    # Try CSV/text if Excel parsing failed
    if not excel_ok:
        # Smart encoding detection: check null byte ratio for UTF-16
        text = None
        null_ratio = sum(1 for b in content if b == 0) / max(len(content), 1)
        encodings = (
            ["utf-16-le", "utf-16-be", "utf-8-sig", "gbk", "gb2312", "gb18030", "utf-8"]
            if null_ratio > 0.15
            else ["utf-8-sig", "gbk", "gb2312", "gb18030", "utf-16-le", "utf-16-be", "utf-8"]
        )
        for enc in encodings:
            try:
                decoded = content.decode(enc)
                sample = decoded[:200]
                printable_ratio = sum(1 for ch in sample if ch.isprintable() or ch in "\t\n\r") / max(len(sample), 1)
                if printable_ratio > 0.7:
                    text = decoded
                    break
            except:
                continue
        if text is None:
            text = content.decode("utf-8", errors="replace")
        text = text.replace("\x00", "")
        delim = "\t" if text.count("\t") > text.count(",") else ","
        try:
            reader = csv.reader(io.StringIO(text), delimiter=delim)
            all_rows = list(reader)
        except:
            for line in text.strip().split("\n"):
                all_rows.append([c.strip() for c in line.split(delim)])

    if not all_rows:
        raise HTTPException(status_code=400, detail="文件为空或无法识别格式")

    # Garbled content check: if cells contain mostly non-printable/non-CJK chars, encoding is wrong
    sample_cells = " ".join(str(c) for c in (all_rows[0] or []) if c)[:200]
    if sample_cells:
        readable = sum(1 for ch in sample_cells if ch.isprintable() or ch in "\t\n\r ") / len(sample_cells)
        has_cjk = bool(re.search(r"[一-鿿]", sample_cells))
        has_ascii = bool(re.search(r"[a-zA-Z]", sample_cells))
        if readable < 0.6 or (not has_cjk and not has_ascii and len(sample_cells) > 10):
            raise HTTPException(
                status_code=400,
                detail=f"文件编码不正确，解析后为乱码。请用 Excel 打开 → 另存为 → CSV UTF-8 格式 → 再导入。当前解析内容: {sample_cells[:100]}..."
            )

    # ── Smart header detection ──
    # Known header aliases (lowercase, no spaces) -> field name
    HEADER_MAP = {
        "level": "level", "lv": "level", "层级": "level", "层次": "level",
        "part_number": "part_number", "partnumber": "part_number", "pn": "part_number",
        "物料号": "part_number", "物料编码": "part_number", "料号": "part_number", "编号": "part_number",
        "name": "name", "名称": "name", "物料名称": "name", "品名": "name", "描述": "name", "description": "name",
        "specification": "spec", "规格": "spec", "型号": "spec", "spec": "spec", "规格型号": "spec",
        "material": "material", "材质": "material", "材料": "material",
        "quantity": "qty", "qty": "qty", "数量": "qty", "用量": "qty", "quantity": "qty",
        "unit": "unit", "单位": "unit",
        "supplier": "supplier", "供应商": "supplier", "厂家": "supplier", "品牌": "supplier",
        "parent": "parent", "parent_part": "parent", "父物料": "parent", "父物料号": "parent",
        "上级": "parent", "所属": "parent", "父级": "parent",
        "备注": "notes", "notes": "notes", "remark": "notes", "说明": "notes",
    }

    def _normalize(s):
        return re.sub(r"[\s_\-（）()]+", "", str(s or "")).lower()

    # Detect if first row is header (contains text) or data
    first_row = all_rows[0]
    text_cells = sum(1 for c in first_row if c and not str(c).strip().isdigit())
    has_header = text_cells >= 2  # at least 2 text cells = header row

    if has_header:
        headers = [_normalize(c) for c in first_row]
        data_start = 1
    else:
        # No header — try position-based fallback
        headers = ["层级", "物料号", "名称", "规格", "材质", "数量", "单位", "供应商", "父物料号"]
        data_start = 0

    # Build column index map
    col_map = {}  # field_name -> column_index
    for idx, h in enumerate(headers):
        mapped = HEADER_MAP.get(h)
        if mapped and mapped not in col_map:
            col_map[mapped] = idx

    # Warn if we couldn't find 'name' column
    if "name" not in col_map:
        # Fallback: assume first non-mapped text column is name
        for idx in range(len(headers)):
            if idx not in col_map.values():
                col_map["name"] = idx
                break

    created = 0
    id_map = {}

    def _get(row, field, default=""):
        idx = col_map.get(field)
        if idx is None or idx >= len(row):
            return default
        v = row[idx]
        return str(v).strip() if v is not None else default

    for i in range(data_start, len(all_rows)):
        row = all_rows[i]
        if not row or not any(c for c in row if c is not None and str(c).strip()):
            continue

        name = _get(row, "name")
        if not name:
            continue

        level_str = _get(row, "level", "0")
        try:
            level = int(float(level_str))
        except ValueError:
            level = 0

        part_number = _get(row, "part_number")
        spec = _get(row, "spec")
        material = _get(row, "material")
        qty_str = _get(row, "qty", "1")
        try:
            qty = float(qty_str)
        except ValueError:
            qty = 1
        unit = _get(row, "unit", "pcs")
        supplier = _get(row, "supplier")
        parent_ref = _get(row, "parent")
        notes = _get(row, "notes")

        parent_id = None
        if parent_ref and parent_ref in id_map:
            parent_id = id_map[parent_ref]
        elif parent_ref:
            # Try to find parent by part_number in DB
            existing = db.execute(
                select(BomItem).where(
                    BomItem.project_id == project_id,
                    BomItem.part_number == parent_ref,
                )
            ).scalar_one_or_none()
            if existing:
                parent_id = existing.id
                id_map[parent_ref] = existing.id

        item = BomItem(project_id=project_id, build_id=build_id, level=level, part_number=part_number,
                       name=name, specification=spec, material=material,
                       quantity=qty, unit=unit or "pcs", supplier=supplier,
                       parent_id=parent_id, sort_order=i, notes=notes)
        db.add(item)
        db.flush()
        id_map[str(i)] = item.id
        if part_number:
            id_map[part_number] = item.id
        created += 1

    db.commit()
    return {"message": f"导入完成", "created": created, "columns_found": list(col_map.keys())}


@router.get("/api/projects/{project_id}/bom/export")
def export_bom(project_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    """Export BOM as Excel download."""
    result = db.execute(select(BomItem).where(BomItem.project_id == project_id).order_by(BomItem.sort_order, BomItem.id))
    items = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(["层级", "物料号", "名称", "规格", "材质", "数量", "单位", "供应商", "父物料号"])
    for item in items:
        ws.append([item.level, item.part_number, item.name, item.specification, item.material,
                    item.quantity, item.unit, item.supplier, ""])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    from fastapi.responses import StreamingResponse
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=BOM-{project_id}.xlsx"})


# ═══════════════════ BOM File Attachments (simple upload, no parsing) ═══════════════════

import uuid, aiofiles
from backend_v2.config import settings

BOM_FILES_DIR = os.path.join(settings.UPLOAD_DIR, "bom_files")
os.makedirs(BOM_FILES_DIR, exist_ok=True)


def _bom_files_meta_path(build_id):
    return os.path.join(BOM_FILES_DIR, f"meta_{build_id}.json")


def _load_bom_files_meta(build_id):
    return safe_load(_bom_files_meta_path(build_id), [])

def _save_bom_files_meta(build_id, data):
    safe_save(_bom_files_meta_path(build_id), data)


@router.post("/api/prototype-builds/{build_id}/bom-files")
async def upload_bom_file(build_id: int, file: UploadFile = File(...),
                          _user=Depends(get_current_user)):
    """Upload a BOM file as attachment to a prototype build."""
    file_id = uuid.uuid4().hex[:12]
    safe_name = f"{build_id}_{file_id}_{file.filename}"
    file_path = os.path.join(BOM_FILES_DIR, safe_name)
    async with aiofiles.open(file_path, "wb") as f:
        content = await file.read()
        await f.write(content)

    meta = _load_bom_files_meta(build_id)
    meta.append({
        "id": file_id,
        "name": file.filename,
        "path": safe_name,
        "size": len(content),
        "size_kb": round(len(content) / 1024, 1),
        "upload_date": datetime.now().strftime("%Y-%m-%d"),
    })
    _save_bom_files_meta(build_id, meta)
    return {"ok": True, "id": file_id}


@router.get("/api/prototype-builds/{build_id}/bom-files")
def list_bom_files(build_id: int, _user=Depends(get_current_user)):
    """List BOM files attached to a prototype build."""
    return _load_bom_files_meta(build_id)


@router.get("/api/prototype-builds/{build_id}/bom-files/{file_id}/download")
def download_bom_file(build_id: int, file_id: str):
    """Download a BOM file attachment."""
    meta = _load_bom_files_meta(build_id)
    entry = next((x for x in meta if x["id"] == file_id), None)
    if not entry:
        raise HTTPException(status_code=404, detail="文件不存在")
    fp = os.path.join(BOM_FILES_DIR, entry["path"])
    if not os.path.exists(fp):
        raise HTTPException(status_code=404, detail="文件不存在")
    return FileResponse(fp, filename=entry["name"])


@router.get("/api/prototype-builds/{build_id}/bom-files/{file_id}/preview")

def preview_bom_file(build_id: int, file_id: str):
    """Preview BOM file content (Excel -> HTML table, text -> pre)."""
    meta = _load_bom_files_meta(build_id)
    entry = next((x for x in meta if x["id"] == file_id), None)
    if not entry:
        raise HTTPException(status_code=404, detail="文件不存在")
    fp = os.path.join(BOM_FILES_DIR, entry["path"])
    if not os.path.exists(fp):
        raise HTTPException(status_code=404, detail="文件不存在")

    fname = entry["name"].lower()

    def _readable(text):
        if not text or len(text.strip()) < 5:
            return False
        # Count truly readable chars: ASCII, common CJK, punctuation
        good = 0
        bad = 0
        for c in text:
            cp = ord(c)
            if c == '�' or cp == 0:
                bad += 1  # replacement char or null
            elif c.isascii():
                good += 1
            elif 0x4e00 <= cp <= 0x9fff:  # CJK Unified
                good += 1
            elif 0x3000 <= cp <= 0x303f:  # CJK punctuation
                good += 1
            elif 0xff00 <= cp <= 0xffef:  # Fullwidth forms
                good += 1
            elif 0x3400 <= cp <= 0x4dbf:  # CJK Extension A
                good += 1
            elif 0xf900 <= cp <= 0xfaff:  # CJK Compatibility
                good += 1
            elif cp <= 0x1f:  # control chars
                bad += 1
            else:
                bad += 1  # unknown script = garbled
        total = good + bad
        if total == 0:
            return False
        if bad / total > 0.15:  # more than 15% garbage
            return False
        return good >= 5

    # Try Excel preview
    if fname.endswith((".xlsx", ".xls")):
        try:
            wb = openpyxl.load_workbook(fp)
            ws = wb.active
            html = '<table style="border-collapse:collapse;width:100%;font-size:12px">'
            all_text = ""
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                html += '<tr>'
                tag = 'th' if i == 0 else 'td'
                for cell in row:
                    v = str(cell or "")
                    all_text += v
                    v = v.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    html += f'<{tag} style="border:1px solid #ddd;padding:4px 8px">{v}</{tag}>'
                html += '</tr>'
            html += '</table>'
            if all_text.strip() and _readable(all_text):
                return {"html": html, "type": "excel"}
        except Exception:
            pass

    # Try text preview with encoding detection
    for enc in ["utf-8-sig", "utf-8", "gbk", "gb2312", "gb18030", "utf-16-le", "utf-16-be"]:
        try:
            with open(fp, "r", encoding=enc) as f:
                text = f.read(50000)
            if text.strip() and _readable(text):
                escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                return {"text": escaped, "type": "text"}
        except Exception:
            continue

    return {"error": "无法预览此文件。建议下载后用 Excel 打开查看。"}

@router.delete("/api/prototype-builds/{build_id}/bom-files/{file_id}")
def delete_bom_file(build_id: int, file_id: str, _user=Depends(get_current_user)):
    meta = _load_bom_files_meta(build_id)
    entry = next((x for x in meta if x["id"] == file_id), None)
    if not entry:
        raise HTTPException(status_code=404, detail="文件不存在")
    fp = os.path.join(BOM_FILES_DIR, entry["path"])
    if os.path.exists(fp):
        try:
            os.remove(fp)
        except:
            pass
    meta = [x for x in meta if x["id"] != file_id]
    _save_bom_files_meta(build_id, meta)
    return {"ok": True}
