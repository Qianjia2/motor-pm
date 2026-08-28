# -*- coding: utf-8 -*-
"""诊断 change 3 附件 PDF 转换丢内容的原因:sheet 结构 / 打印区域 / 分页。"""
import io, os, sys, subprocess, tempfile, shutil
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import requests

BASE = "http://localhost:5002"
s = requests.Session()
tok = s.post(f"{BASE}/api/auth/login", json={"username": "admin", "password": "admin123"}).json()["access_token"]
h = {"Authorization": f"Bearer {tok}"}
r = s.get(f"{BASE}/api/changes/3", headers=h)
print("changes/3 status:", r.status_code)
raw = s.get(f"{BASE}/api/changes/3/attachment", headers=h).content
print("raw bytes:", len(raw))
open("_diag_orig.xlsx", "wb").write(raw)

import openpyxl
wb = openpyxl.load_workbook("_diag_orig.xlsx", data_only=True)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    psp = ws.sheet_properties.pageSetUpPr
    print(f"  [{ws.title}] dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} "
          f"merged={len(ws.merged_cells.ranges)} print_area={ws.print_area!r} state={ws.sheet_state}")
    print(f"    pageSetup: orient={ws.page_setup.orientation} fitToPage={psp.fitToPage if psp else None} "
          f"scale={ws.page_setup.scale} paper={ws.page_setup.paperSize}")
    print(f"    images={len(ws._images)} breaks={list(ws.row_breaks.brk)}")

# 现有转换:soffice 转 PDF 并逐页看末页文本
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
tmp = tempfile.mkdtemp(prefix="lo_diag_")
try:
    subprocess.run([SOFFICE, "--headless", "--norestore", "--convert-to", "pdf",
                    "--outdir", tmp, "_diag_orig.xlsx"], capture_output=True, timeout=120)
    try:
        from pypdf import PdfReader
    except ImportError:
        from PyPDF2 import PdfReader
    rd = PdfReader(os.path.join(tmp, "_diag_orig.pdf"))
    print("PDF pages:", len(rd.pages))
    for i, pg in enumerate(rd.pages, 1):
        t = (pg.extract_text() or "").strip()
        print(f"--- page {i} ({len(t)} chars): {t[:150].replace(chr(10), ' | ')}")
    # 关键单元格是否出现在 PDF
    full = " ".join((pg.extract_text() or "") for pg in rd.pages)
    for probe in ["怎么填", "Easi-ECR", "变更前内容", "变更后内容", "处置", "确认本学科影响可接受"]:
        print(f"probe {probe!r}: {'FOUND' if probe in full else 'MISSING'}")
finally:
    shutil.rmtree(tmp, ignore_errors=True)
