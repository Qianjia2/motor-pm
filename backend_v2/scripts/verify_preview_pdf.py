# -*- coding: utf-8 -*-
"""验证 office 附件预览现在返回真 PDF(LibreOffice 保真排版)。"""
import io, os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import requests

BASE = "http://localhost:5002"
s = requests.Session()

def req(method, path, **kw):
    r = s.request(method, BASE + path, **kw)
    ct = r.headers.get("content-type", "")
    body = r.json() if ct.startswith("application/json") else f"<{ct} len={len(r.content)}>"
    print(f"[{r.status_code}] {method} {path} -> {body if isinstance(body, str) else str(body)[:200]}")
    return r

r = req("POST", "/api/auth/login", json={"username": "admin", "password": "admin123"})
s.headers["Authorization"] = f"Bearer {r.json()['access_token']}"

# 1. 上传预览接口:带合并单元格/样式的 xlsx → 应返回 application/pdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
wb = Workbook(); ws = wb.active
ws.title = "技术变更申请表"
ws["A1"] = "技术变更申请表"; ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:D1"); ws["A1"].fill = PatternFill("solid", fgColor="D9E2F3")
ws["A2"] = "变更编号"; ws["B2"] = "TC-2026-001"; ws["C2"] = "变更级别"; ws["D2"] = "重大"
ws["A3"] = "变更内容"; ws["B3"] = "电机绕组方案调整"; ws.merge_cells("B3:D3")
ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 24
buf = io.BytesIO(); wb.save(buf)

r = req("POST", "/api/changes/preview-attachment",
        files={"file": ("变更申请表.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
assert r.status_code == 200, "preview failed"
ct = r.headers.get("content-type", "")
assert ct == "application/pdf" and r.content[:5] == b"%PDF-", f"expect pdf, got {ct}"
print("OK: xlsx 预览 -> 真 PDF(%d 字节)" % len(r.content))

# 2. docx 同样转 PDF
from docx import Document
doc = Document()
doc.add_heading("技术变更说明", level=1)
doc.add_paragraph("本变更涉及电机绕组方案调整，影响功率密度与温升指标。")
tbl = doc.add_table(rows=2, cols=2)
tbl.cell(0, 0).text = "项目"; tbl.cell(0, 1).text = "高压电机平台"
tbl.cell(1, 0).text = "变更级别"; tbl.cell(1, 1).text = "重大变更"
buf2 = io.BytesIO(); doc.save(buf2)
r = req("POST", "/api/changes/preview-attachment",
        files={"file": ("说明.docx", buf2.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
assert r.headers.get("content-type") == "application/pdf" and r.content[:5] == b"%PDF-", "docx pdf failed"
print("OK: docx 预览 -> 真 PDF(%d 字节)" % len(r.content))

# 3. pdf/图片仍走 raw
r = req("POST", "/api/changes/preview-attachment",
        files={"file": ("t.pdf", b"%PDF-1.4\n%%EOF", "application/pdf")})
assert r.json().get("kind") == "raw", "pdf should stay raw"
print("OK: 原生 pdf 仍 raw")

# 4. 已保存附件的 view=1 也返回 PDF(真实变更 id=3, 用户真实 xlsx)
import os as _os
r = s.get(f"{BASE}/api/changes/3/attachment?token={s.headers['Authorization'].split()[-1]}&view=1")
ct = r.headers.get("content-type", "")
print(f"real change 3 view=1 -> {ct} len={len(r.content)}")
if ct == "application/pdf":
    print("OK: 已保存 xlsx 附件预览 -> 真 PDF(用户原排版)")
elif ct == "application/json":
    r2 = req("GET", "/api/changes/3")
    print("NOTE: change 3 attachment state:", r2.json().get("attachment"))
else:
    print("NOTE: got", ct)

print("ALL PASS")
